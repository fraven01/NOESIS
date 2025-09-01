from django.contrib.auth.models import User, Group
from django.urls import reverse
from django.conf import settings
from django.http import QueryDict
from django.db import IntegrityError
from types import SimpleNamespace
import os
import re
import pytest

pytestmark = pytest.mark.unit

from django.apps import apps
from ...models import (
    BVProject,
    BVProjectFile,
    Recording,
    Prompt,
    LLMConfig,
    Tile,
    UserTileAccess,
    GroupTileAccess,
    Anlage1Question,
    Anlage1Config,
    Area,
    Anlage2Function,
    Anlage2Config,
    Anlage2ColumnHeading,
    Anlage2SubQuestion,
    AnlagenFunktionsMetadaten,
    FunktionsErgebnis,
    SoftwareKnowledge,
    BVSoftware,
    Gutachten,
    AntwortErkennungsRegel,
    Anlage4Config,
    Anlage4ParserConfig,
    ZweckKategorieA,
    Anlage5Review,
)
from ...docx_utils import (
    extract_text,
    get_docx_page_count,
    get_pdf_page_count,
    parse_anlage2_table,
    _normalize_header_text,
)

from ...utils import start_analysis_for_file
from ... import text_parser

from core.text_parser import parse_anlage2_text, PHRASE_TYPE_CHOICES

from ...anlage4_parser import parse_anlage4

from ...parser_manager import parser_manager
from ...parsers import AbstractParser

from pathlib import Path
from tempfile import NamedTemporaryFile, TemporaryDirectory
from io import BytesIO
from docx import Document
import shutil
from PIL import Image
import fitz

from django.core.files.uploadedfile import SimpleUploadedFile
from ...forms import (
    BVProjectForm,
    BVProjectUploadForm,
    BVProjectFileJSONForm,
    BVProjectFileForm,
    Anlage2ConfigForm,
    Anlage2ReviewForm,
)
from ...workflow import set_project_status
from ...models import ProjectStatus
from ...llm_tasks import (
    check_anlage1,
    check_anlage2,
    analyse_anlage3,
    analyse_anlage4,
    analyse_anlage4_async,
    check_anlage5,
    run_conditional_anlage2_check,
    worker_verify_feature,
    worker_generate_gutachten,
    worker_run_initial_check,
    worker_anlage4_evaluate,
    worker_generate_gap_summary,
    summarize_anlage1_gaps,
    summarize_anlage2_gaps,
    get_prompt,
    generate_gutachten,
    run_anlage2_analysis,
    parse_anlage1_questions,
)
from ...views import (
    _verification_to_initial,
    _build_row_data,
    _build_supervision_row,
    _save_project_file,
    extract_anlage_nr,
    get_user_tiles,
    _has_manual_gap,
    _build_supervision_groups,
    _resolve_value,
)
from ...reporting import generate_gap_analysis, generate_management_summary
from unittest.mock import patch, ANY, Mock, call
from django.core.management import call_command
from django.test import override_settings
import json
from ..base import NoesisTestCase
from ...initial_data_constants import INITIAL_PROJECT_STATUSES
from ...prompt_context import build_prompt_context, available_placeholders
from ..utils import (
    create_project,
    seed_test_data,
    DEFAULT_STATUS_KEY,
    _extra_statuses,
)





def test_build_prompt_context_keys(db) -> None:
    """Prüft, dass ``build_prompt_context`` alle Platzhalter füllt."""

    projekt = create_project(title="Demo", software=["Soft"])
    ctx = build_prompt_context(projekt)
    for key in available_placeholders():
        assert key in ctx
    assert ctx["project_name"] == "Demo"
    assert ctx["software_name"] == "Soft"

    Prompt.objects.update_or_create(
        name="check_anlage3_vision",
        defaults={
            "text": "Prüfe die folgenden Bilder der Anlage. Gib ein JSON mit 'ok' und 'hinweis' zurück:\n\n"
        },
    )
    # Weitere Prompts für Tests bereitstellen
    roles = {r.name: r for r in apps.get_model("core", "LLMRole").objects.all()}
    prompt_data = {
        "anlage1_email": {
            "text": (
                "Formuliere eine freundliche E-Mail an den Fachbereich. "
                "Wir haben die Anlage 1 geprüft und noch folgende Vorschläge, "
                "bevor der Mitbestimmungsprozess weiter gehen kann. "
                "Bitte fasse die folgenden Vorschläge zusammen:\r\n\r\n"
            )
        },
        "anlage2_ai_involvement_check": {
            "text": (
                "Antworte ausschließlich mit 'Ja' oder 'Nein'. Frage: "
                "Beinhaltet die Funktion '{function_name}' der Software "
                "'{software_name}' typischerweise eine KI-Komponente? "
                "Eine KI-Komponente liegt vor, wenn die Funktion "
                "unstrukturierte Daten (Text, Bild, Ton) verarbeitet, "
                "Sentiment-Analysen durchführt oder nicht-deterministische, "
                "probabilistische Ergebnisse liefert."
            )
        },
        "anlage2_feature_justification": {
            "text": (
                "Warum besitzt die Software '{software_name}' typischerweise die "
                "Funktion oder Eigenschaft '{function_name}'?   Ist es möglich "
                "mit der {function_name} eine Leistungskontrolle oder eine "
                "Verhaltenskontrolle  Ist damit eine Leistungskontrolle oder "
                "eine Verhaltenskontrolle im Sinne des §87 Abs. 1 Nr. 6 möglich? "
                "Wenn ja, wie?"
            )
        },
        "anlage2_subquestion_possibility_check": {
            "text": (
                "Im Kontext der Funktion '{function_name}' der Software "
                "'{software_name}': Ist die spezifische Anforderung "
                "'{subquestion_text}' technisch möglich? Antworte nur mit 'Ja', "
                "'Nein' oder 'Unsicher'."
            )
        },
        "anlage2_subquestion_justification_check": {
            "text": (
                " [SYSTEM]\nDu bist Fachautor*in für IT-Mitbestimmung (§87 Abs. 1 Nr. 6 BetrVG).\n"
                "Antworte Unterfrage prägnant in **maximal zwei Sätzen** (insgesamt ≤ 65 Wörter) und erfülle folgende Regeln :\n\n"
                "1. Starte Teil A mit „Typischer Zweck: …“  \n2. Starte Teil B mit „Kontrolle: Ja, …“ oder „Kontrolle: Nein, …“.  \n"
                "3. Nenne exakt die übergebene Funktion/Eigenschaft, erfinde nichts dazu.  \n"
                "4. Erkläre knapp *warum* mit der Funktion die Unterfrage (oder warum nicht) eine Leistungs- oder Verhaltenskontrolle möglich ist.  \n"
                "5. Verwende Alltagssprache, keine Marketing-Floskeln.\n\n"
                ' [USER]\nSoftware: {{software_name}}  \nFunktion/Eigenschaft: {{function_name}}  \nUnterfrage: "{{subquestion_text}}"'
            )
        },
        "anlage2_ai_verification_prompt": {
            "text": (
                "Gib eine kurze Begründung, warum die Funktion '{function_name}' "
                "(oder die Unterfrage '{subquestion_text}') der Software "
                "'{software_name}' eine KI-Komponente beinhaltet oder beinhalten kann, "
                "insbesondere im Hinblick auf die Verarbeitung unstrukturierter Daten "
                "oder nicht-deterministischer Ergebnisse."
            )
        },
        "anlage2_feature_verification": {
            "text": (
                "Deine einzige Aufgabe ist es, die folgende Frage mit einem einzigen "
                'Wort zu beantworten. Deine Antwort darf AUSSCHLIESSLICH "Ja", '
                '"Nein", oder "Unsicher" sein. Gib keine Einleitung, keine '
                "Begründung und keine weiteren Erklärungen ab.\r\n\r\nFrage: "
                "Besitzt die Software '{software_name}' basierend auf allgemeinem "
                "Wissen typischerweise die Funktion oder Eigenschaft "
                "'{function_name}'?\n\n{gutachten}"
            ),
            "role": roles.get("Standard"),
            "use_system_role": False,
        },
        "check_anlage2_function": {
            "text": (
                "Prüfe anhand des folgenden Textes, ob die genannte Funktion "
                'vorhanden ist. Gib ein JSON mit den Schlüsseln "technisch_verfuegbar", '
                '"einsatz_telefonica", "zur_lv_kontrolle" und "ki_beteiligung" '
                "zurück.\n\n"
            )
        },
        "check_anlage4": {
            "text": "Prüfe die folgende Anlage auf Vollständigkeit. Gib ein JSON mit 'ok' und 'hinweis' zurück:\n\n",
        },
        "check_anlage5": {
            "text": "Prüfe die folgende Anlage auf Vollständigkeit. Gib ein JSON mit 'ok' und 'hinweis' zurück:\n\n"
        },
        "classify_system": {
            "text": (
                "Bitte klassifiziere das folgende Softwaresystem. "
                "Gib ein JSON mit den Schlüsseln 'kategorie' und 'begruendung' zurück.\n\n"
            )
        },
        "generate_gutachten": {
            "text": (
                "Erstelle ein tiefgehendes Gutachten zu der Software im Sinne des § 87 Abs. 1 Nr. 6 BetrVG. "
                "Richte das Gutachten ausschließlich an Betriebsräte und überspringe allgemeine Erläuterungen "
                "zu DSGVO oder Datenschutzrecht ebenso musst du nicht erläutern, wann Mitbestimmungsrechte "
                "nach §87 (1) Nr. 6 gelten.\n\nDein Gutachten soll folgende Punkte abdecken: \n\n1. **Mitbestimmungspflichtige Funktionen**   \n- Liste alleFeatures auf, die der Leistungs- oder Verhaltenskontrolle dienen (z. B. Analyse von Nutzungshistorien, App- und Kommunikationsauswertung, Dateizugriffsprotokolle).\n- Erläutere kurz, warum jede einzelne Funktion unter § 87 1 Nr. 6 BetrVG fällt. \n\n2. **Eingriffsintensität aus Mitarbeitersicht**   \n   - Beschreibe, wie stark jede dieser Funktionen in den Arbeitsablauf eingreift und welche Verhaltensaspekte sie überwacht (z. B. Häufigkeit von App-Nutzung, Kommunikationsverhalten, Standortdaten).   \n   - Nutze eine Skala (gering – mittel – hoch) und begründe die Einstufung anhand typischer Betriebsabläufe in einem Telekommunikationsunternehmen. \n\n3. **Betroffene Leistungs- und Verhaltensaspekte**   \n   - Identifiziere konkret, welche Leistungskennzahlen (z. B. Aktivitätszeiten, App-Nutzungsdauer) und Verhaltensmuster (z. B. Kommunikationshäufigkeit, Datenübertragung) erfasst und ausgewertet werden.   \n   - Schätze ab, wie umfassend und detailliert die Auswertung jeweils ausfällt. \n\n4. **Handlungsbedarf für den Betriebsrat**   \n   - Fasse zusammen, bei welchen Funktionen und Einsatzszenarien eine Betriebsvereinbarung zwingend erforderlich ist. \n\n5. **Weitere Mitbestimmungsrechte (Kurzhinweise)**   \n   - Wenn offensichtlich erkennbar ist, dass andere relevante Mitbestimmungsrechte nach BetrVG (z. B. §§ 80 ff. zur Informationspflicht) berührt sind, bewerte kurz, warum diese Software dieses Recht des Betriebsrats berühren könnte. \n\nArbeite strukturiert mit klaren Überschriften und Bullet-Points. Wo sinnvoll, nutze kurze Tabellen oder Zusammenfassungen zur Übersicht. \n. Antworte auf deutsch.\nSoftware: \n"
            ),
            "role": roles.get("Gutachten"),
        },
        "initial_check_knowledge": {
            "text": "Kennst du die Software '{name}'? Antworte ausschließlich mit einem einzigen Wort: 'Ja' oder 'Nein'.",
            "use_system_role": False,
        },
        
        "initial_llm_check": {
            "text": (
                "Erstelle eine kurze, technisch korrekte Beschreibung für die Software '{name}'. "
                "Nutze Markdown mit Überschriften, Listen oder Fettdruck, um den Text zu strukturieren. "
                "Erläutere, was sie tut und wie sie typischerweise eingesetzt wird."
            ),
            "role": roles.get("Gutachten"),
        },
    }

    for name, data in prompt_data.items():
        Prompt.objects.update_or_create(
            name=name,
            defaults={
                "text": data["text"],
                "role": data.get("role"),
                "use_system_role": data.get("use_system_role", True),
            },
        )


@pytest.mark.usefixtures("seed_db")
class SeedInitialDataTests(NoesisTestCase):
    """Stellt sicher, dass die Seed-Daten vorhanden sind."""

    def test_answer_rules_seeded(self) -> None:
        """Prüft die durch die globale Fixture angelegten Antwortregeln."""
        from ...initial_data_constants import INITIAL_ANSWER_RULES

        for rule in INITIAL_ANSWER_RULES:
            obj = AntwortErkennungsRegel.objects.get(
                regel_name=rule["regel_name"]
            )
            self.assertEqual(
                obj.erkennungs_phrase,
                rule["erkennungs_phrase"],
            )
            self.assertEqual(
                obj.actions_json,
                rule["actions"],
            )


class ExtractAnlageNrTests(NoesisTestCase):
    """Tests für die Erkennung der Anlagen-Nummer aus Dateinamen."""

    def test_variants(self):
        self.assertEqual(extract_anlage_nr("Anlage 1.docx"), 1)
        self.assertEqual(extract_anlage_nr("Anlage-2.pdf"), 2)
        self.assertEqual(extract_anlage_nr("Anlage3.docx"), 3)


@pytest.mark.usefixtures("seed_db")
class BVProjectFileTests(NoesisTestCase):
    def setUp(self) -> None:  # pragma: no cover - setup
        super().setUp()
        self.anmelden_func = Anlage2Function.objects.create(name="Anmelden")
        self.superuser = User.objects.get(username="frank")

    def test_create_project_with_files(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        for i in range(1, 4):
            f = SimpleUploadedFile(f"f{i}.txt", b"data")
            BVProjectFile.objects.create(
                project=projekt,
                anlage_nr=i,
                upload=f,
                text_content="data",
            )
        status_before = projekt.status.key
        self.assertEqual(projekt.anlagen.count(), 3)
        self.assertListEqual(
            list(projekt.anlagen.values_list("anlage_nr", flat=True)), [1, 2, 3]
        )

    def test_default_flags(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("a.txt", b"data"),
            text_content="d",
        )
        self.assertFalse(pf.manual_reviewed)
        self.assertFalse(pf.verhandlungsfaehig)
        pf.manual_reviewed = True
        pf.verhandlungsfaehig = True
        pf.save()
        pf.refresh_from_db()
        self.assertTrue(pf.manual_reviewed)
        self.assertTrue(pf.verhandlungsfaehig)

    @pytest.mark.slow
    def test_project_delete_removes_files(self):
        """Sichert, dass beim Löschen eines Projekts die Dateien entfernt werden."""
        with TemporaryDirectory() as tmpdir, override_settings(MEDIA_ROOT=tmpdir):
            projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
            pf = BVProjectFile.objects.create(
                project=projekt,
                anlage_nr=1,
                upload=SimpleUploadedFile("a.txt", b"data"),
                text_content="d",
            )
            file_path = pf.upload.path
            self.assertTrue(os.path.exists(file_path))
            projekt.delete()
            self.assertFalse(os.path.exists(file_path))

    def test_auto_start_analysis_saves_task_id(self):
        """Die Analyse-ID wird nach dem Upload gespeichert."""
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        with patch(
            "core.signals.start_analysis_for_file", return_value="tid"
        ) as mock_start:
            pf = BVProjectFile.objects.create(
                project=projekt,
                anlage_nr=1,
                upload=SimpleUploadedFile("a.txt", b"data"),
            )
        mock_start.assert_called_with(pf.pk)
        pf.refresh_from_db()
        self.assertEqual(pf.verification_task_id, "tid")

    def test_json_form_shows_analysis_field_for_anlage3(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=3,
            upload=SimpleUploadedFile("a.txt", b"x"),
            text_content="t",
            analysis_json={"ok": True},
        )
        form = BVProjectFileJSONForm(instance=pf)
        self.assertIn("analysis_json", form.fields)

    def test_save_does_not_start_task(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        with patch("core.models.async_task") as mock_task:
            pf = BVProjectFile.objects.create(
                project=projekt,
                anlage_nr=2,
                upload=SimpleUploadedFile("a.txt", b"x"),
                text_content="t",
            )
        self.assertEqual(pf.verification_task_id, "")
        mock_task.assert_not_called()

    def test_is_verification_running(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"x"),
            verification_task_id="tid",
        )
        with patch("core.models.fetch") as mock_fetch:
            mock_fetch.return_value = SimpleNamespace(success=None)
            self.assertTrue(pf.is_verification_running())
            mock_fetch.return_value = SimpleNamespace(success=True)
            self.assertFalse(pf.is_verification_running())

    def test_check_functions_clears_task_id(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        with patch("core.signals.start_analysis_for_file", return_value="tid"):
            pf = BVProjectFile.objects.create(
                project=projekt,
                anlage_nr=2,
                upload=SimpleUploadedFile("a.txt", b"x"),
                verification_task_id="tid",
            )
        _ = self.anmelden_func
        with (
            patch("core.llm_tasks.query_llm", return_value="{}"),
            patch("core.llm_tasks.async_task") as mock_async,
            patch("core.llm_tasks.result") as mock_result,
        ):
            mock_async.side_effect = lambda name, *a, **k: (
                worker_verify_feature(*a, **k) or "tid"
            )
            mock_result.side_effect = lambda *a, **k: None
            run_conditional_anlage2_check(pf.pk)
        pf.refresh_from_db()
        self.assertEqual(pf.verification_task_id, "")
        self.assertEqual(pf.processing_status, BVProjectFile.COMPLETE)

    def test_template_shows_disabled_state_when_task_running(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"x"),
            verification_task_id="tid",
        )
        self.client.login(username=self.superuser.username, password="pass")
        with patch("core.models.fetch") as mock_fetch:
            mock_fetch.return_value = SimpleNamespace(success=None)
            url = reverse("projekt_detail", args=[projekt.pk])
            resp = self.client.get(url)
        self.assertContains(resp, "animate-spin")  # Spinner-Element bei laufender Analyse

    def test_hx_project_file_status_running(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        with patch("core.signals.start_analysis_for_file", return_value="tid"):
            pf = BVProjectFile.objects.create(
                project=projekt,
                anlage_nr=2,
                upload=SimpleUploadedFile("a.txt", b"x"),
                verification_task_id="tid",
            )
        self.client.login(username=self.superuser.username, password="pass")
        with patch("core.models.fetch") as mock_fetch:
            mock_fetch.return_value = SimpleNamespace(success=None)
            url = reverse("hx_anlage_status", args=[pf.pk])
            resp = self.client.get(url)
        self.assertContains(resp, "hx-trigger")
        self.assertContains(resp, "animate-spin")

    def test_hx_project_file_status_ready(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        with patch("core.signals.start_analysis_for_file", return_value="tid"):
            pf = BVProjectFile.objects.create(
                project=projekt,
                anlage_nr=2,
                upload=SimpleUploadedFile("a.txt", b"x"),
                verification_task_id="tid",
            )
        self.client.login(username=self.superuser.username, password="pass")
        url = reverse("hx_anlage_status", args=[pf.pk])
        with patch("core.models.fetch") as mock_fetch:
            mock_fetch.return_value = SimpleNamespace(success=True)
            resp = self.client.get(url)
        self.assertContains(resp, "hx-trigger")
        # Finalen Status simulieren
        pf.processing_status = BVProjectFile.COMPLETE
        pf.verification_task_id = ""
        pf.save()
        resp = self.client.get(url)
        self.assertNotContains(resp, "hx-trigger")
        self.assertContains(resp, "Analyse bearbeiten")

    def test_hx_anlage_status_processing(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"x"),
            processing_status=BVProjectFile.PROCESSING,
        )
        self.client.login(username=self.superuser.username, password="pass")
        url = reverse("hx_anlage_status", args=[pf.pk])
        resp = self.client.get(url)
        self.assertContains(resp, "animate-spin")

    def test_hx_anlage_status_ready(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        # Der Signal-Handler `auto_start_analysis` speichert die
        # R\xFCckgabe von `start_analysis_for_file` als Task-ID. Wenn der
        # Mock kein einfaches String-Ergebnis liefert, w\xFCrde ein
        # `FieldError` auftreten. Daher liefern wir explizit eine
        # Zeichenkette zur\xFCck.
        with patch(
            "core.signals.start_analysis_for_file", return_value="mocked_task_id"
        ):
            pf = BVProjectFile.objects.create(
                project=projekt,
                anlage_nr=2,
                upload=SimpleUploadedFile("a.txt", b"x"),
                processing_status=BVProjectFile.COMPLETE,
                analysis_json={},
            )
        self.client.login(username=self.superuser.username, password="pass")
        url = reverse("hx_anlage_status", args=[pf.pk])
        resp = self.client.get(url)
        self.assertContains(resp, "Analyse bearbeiten")
        self.assertContains(resp, "Erneut analysieren")

    def test_hx_anlage_status_failed(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        with patch("core.signals.start_analysis_for_file", return_value=""):
            pf = BVProjectFile.objects.create(
                project=projekt,
                anlage_nr=2,
                upload=SimpleUploadedFile("a.txt", b"x"),
                processing_status=BVProjectFile.FAILED,
            )
        self.client.login(username=self.superuser.username, password="pass")
        url = reverse("hx_anlage_status", args=[pf.pk])
        resp = self.client.get(url)
        self.assertNotContains(resp, "hx-trigger")
        self.assertContains(resp, "Analyse fehlgeschlagen")
        self.assertContains(resp, "Erneut versuchen")

    def test_hx_anlage_status_pending(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        with patch("core.signals.start_analysis_for_file", return_value=""):
            pf = BVProjectFile.objects.create(
                project=projekt,
                anlage_nr=2,
                upload=SimpleUploadedFile("a.txt", b"x"),
                processing_status=BVProjectFile.PENDING,
            )
        self.client.login(username=self.superuser.username, password="pass")
        url = reverse("hx_anlage_status", args=[pf.pk])
        resp = self.client.get(url)
        self.assertContains(resp, "Analyse starten")
        self.assertContains(resp, "hx-trigger")

    def test_hx_project_anlage_tab(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        with patch("core.signals.start_analysis_for_file", return_value=""):
            pf = BVProjectFile.objects.create(
                project=projekt,
                anlage_nr=1,
                upload=SimpleUploadedFile("a.txt", b"x"),
            )
        self.client.login(username=self.superuser.username, password="pass")
        url = reverse("hx_project_anlage_tab", args=[projekt.pk, 1])
        resp = self.client.get(url)
        self.assertContains(resp, 'href="/media/bv_files/a_')
        self.assertContains(resp, '.txt"')
        self.assertContains(resp, "hx-trigger")

    def test_hx_anlage_row(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        with patch("core.signals.start_analysis_for_file", return_value=""):
            pf = BVProjectFile.objects.create(
                project=projekt,
                anlage_nr=1,
                upload=SimpleUploadedFile("a.txt", b"x"),
            )
        self.client.login(username=self.superuser.username, password="pass")
        url = reverse("hx_anlage_row", args=[pf.pk])
        resp = self.client.get(url)
        self.assertContains(resp, 'href="/media/bv_files/a_')
        self.assertContains(resp, '.txt"')
        self.assertContains(resp, "hx-trigger")

    def test_hx_toggle_project_file_flag(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("a.txt", b"x"),
        )
        self.client.login(username=self.superuser.username, password="pass")
        url = reverse("hx_toggle_project_file_flag", args=[pf.pk, "manual_reviewed"])
        resp = self.client.post(
            url,
            {"value": "1"},
            HTTP_HX_REQUEST="true",
        )
        self.assertEqual(resp.status_code, 200)
        pf.refresh_from_db()
        self.assertTrue(pf.manual_reviewed)
        self.assertContains(resp, "<tr")
        self.assertContains(resp, "✓")

    def test_toggle_verhandlungsfaehig_sets_project_done(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        for nr in range(1, 6):
            BVProjectFile.objects.create(
                project=projekt,
                anlage_nr=nr,
                upload=SimpleUploadedFile(f"a{nr}.txt", b"x"),
                verhandlungsfaehig=True,
            )
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=6,
            upload=SimpleUploadedFile("a6.txt", b"x"),
            verhandlungsfaehig=False,
        )
        status_before = projekt.status.key
        self.client.login(username=self.superuser.username, password="pass")
        url = reverse("project_file_toggle_flag", args=[pf.pk, "verhandlungsfaehig"])
        resp = self.client.post(url, {"value": "1"})
        self.assertEqual(resp.status_code, 302)
        projekt.refresh_from_db()
        self.assertNotEqual(projekt.status.key, status_before)
        self.assertEqual(projekt.status.key, "DONE")
        self.assertTrue(all(f.verhandlungsfaehig for f in projekt.anlagen.all()))

    def test_hx_project_software_tab(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        SoftwareKnowledge.objects.create(project=projekt, software_name="A")
        self.client.login(username=self.superuser.username, password="pass")
        url = reverse("hx_project_software_tab", args=[projekt.pk, "tech"])
        resp = self.client.get(url)
        self.assertContains(resp, "Prüfung starten")

    def test_trigger_file_analysis_starts_tasks(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("a.txt", b"x"),
            processing_status=BVProjectFile.PENDING,
        )
        self.client.login(username=self.superuser.username, password="pass")
        with patch("core.views.start_analysis_for_file", return_value="123") as mock_start:
            url = reverse("trigger_file_analysis", args=[pf.pk])
            resp = self.client.post(url)
        mock_start.assert_called_with(pf.pk)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.json(), {"task_id": "123"})

    def test_start_analysis_for_file_enqueues_tasks(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("a.txt", b"x"),
            processing_status=BVProjectFile.PENDING,
        )
        with patch.object(
            BVProjectFile,
            "get_analysis_tasks",
            return_value=[("core.llm_tasks.check_anlage1", pf.pk)],
        ), patch("core.utils.async_task") as mock_async, patch(
            "core.utils.transaction.on_commit", side_effect=lambda func: func()
        ):
            mock_async.return_value = "t1"
            task_id = start_analysis_for_file(pf.pk)
        mock_async.assert_called_with("core.llm_tasks.check_anlage1", pf.pk)
        self.assertEqual(task_id, "t1")
        pf.refresh_from_db()
        self.assertEqual(pf.processing_status, BVProjectFile.PROCESSING)

    def test_get_analysis_tasks_returns_project_id_for_conditional_check(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"x"),
        )
        tasks = pf.get_analysis_tasks()
        self.assertEqual(
            tasks,
            [
                ("core.llm_tasks.worker_run_anlage2_analysis", pf.pk),
                ("core.llm_tasks.run_conditional_anlage2_check", pf.pk),
            ],
        )




class BVProjectModelTests(NoesisTestCase):
    def test_title_auto_set_from_software(self):
        projekt = BVProject.objects.create(software_typen="A, B", beschreibung="x")
        self.assertEqual(projekt.title, "A, B")

    def test_title_preserved_when_set(self):
        projekt = BVProject.objects.create(
            title="X", software_typen="A", beschreibung="x"
        )
        self.assertEqual(projekt.title, "X")

    def test_save_accepts_list_for_software_typen(self):
        projekt = BVProject.objects.create(
            software_typen=["A", "", "B "],
            beschreibung="x",
        )
        self.assertEqual(projekt.software_typen, "A, B")


class AnlagenFunktionsMetadatenModelTests(NoesisTestCase):
    def setUp(self) -> None:  # pragma: no cover - setup
        super().setUp()
        self.anmelden_func = Anlage2Function.objects.create(name="Anmelden")

    def test_manual_result_field(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        func = self.anmelden_func
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"x"),
        )
        res = AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf,
            funktion=func,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            quelle="manuell",
            technisch_verfuegbar=True,
            ki_beteiligung=False,
        )
        latest = FunktionsErgebnis.objects.filter(
            anlage_datei__project=projekt,
            funktion=func,
            quelle="manuell",
        ).first()
        self.assertTrue(latest.technisch_verfuegbar)
        self.assertFalse(latest.ki_beteiligung)


class WorkflowTests(NoesisTestCase):
    def test_default_status(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        self.assertEqual(projekt.status.key, DEFAULT_STATUS_KEY)

    def test_set_project_status(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        set_project_status(projekt, "IN_PROGRESS")
        projekt.refresh_from_db()
        self.assertEqual(projekt.status.key, "IN_PROGRESS")

    def test_invalid_status(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        with self.assertRaises(ValueError):
            set_project_status(projekt, "XXX")

    def test_status_history_created(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        self.assertEqual(projekt.status_history.count(), 1)
        set_project_status(projekt, "IN_PROGRESS")
        self.assertEqual(projekt.status_history.count(), 2)


class BuildRowDataTests(NoesisTestCase):
    def setUp(self):
        self.func = Anlage2Function.objects.create(name="Testfunktion")
        self.form = Anlage2ReviewForm()

    def test_flag_set_on_difference(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"x"),
        )
        res = AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf,
            funktion=self.func,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            quelle="parser",
            technisch_verfuegbar=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            quelle="ki",
            technisch_verfuegbar=False,
        )
        result_map = {res.get_lookup_key(): res}

        row = _build_row_data(
            "Testfunktion",
            "Testfunktion",
            self.func.id,
            f"func{self.func.id}_",
            self.form,
            {},
            {},
            {},
            {},
            result_map,
        )
        self.assertTrue(row["requires_manual_review"])

    def test_flag_not_set_when_manual(self):
        manual = {"Testfunktion": {"technisch_vorhanden": True}}
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"x"),
        )
        res = AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf,
            funktion=self.func,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            quelle="parser",
            technisch_verfuegbar=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            quelle="ki",
            technisch_verfuegbar=False,
        )
        result_map = {res.get_lookup_key(): res}

        row = _build_row_data(
            "Testfunktion",
            "Testfunktion",
            self.func.id,
            f"func{self.func.id}_",
            self.form,
            {},
            {},
            {},
            manual,
            result_map,
        )
        self.assertFalse(row["requires_manual_review"])

    def test_manual_flags_propagated(self):
        manual = {"Testfunktion": {"technisch_vorhanden": True}}
        row = _build_row_data(
            "Testfunktion",
            "Testfunktion",
            self.func.id,
            f"func{self.func.id}_",
            self.form,
            {},
            {},
            {},
            manual,
            {},
        )
        self.assertTrue(row["manual_flags"]["technisch_vorhanden"])

    def test_manual_flags_false_when_absent(self):
        row = _build_row_data(
            "Testfunktion",
            "Testfunktion",
            self.func.id,
            f"func{self.func.id}_",
            self.form,
            {},
            {},
            {},
            {},
            {},
        )
        self.assertFalse(row["manual_flags"]["technisch_vorhanden"])

    def test_doc_ai_from_result_map_main(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"x"),
        )
        res = AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf,
            funktion=self.func,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            quelle="parser",
            technisch_verfuegbar=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            quelle="ki",
            technisch_verfuegbar=False,
        )
        result_map = {res.get_lookup_key(): res}

        row = _build_row_data(
            self.func.name,
            self.func.name,
            self.func.id,
            f"func{self.func.id}_",
            self.form,
            {},
            {},
            {},
            {},
            result_map,
        )

        self.assertTrue(row["doc_result"]["technisch_vorhanden"])
        self.assertFalse(row["ai_result"]["technisch_vorhanden"])

    def test_doc_ai_from_result_map_subquestion(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        sub = Anlage2SubQuestion.objects.create(
            funktion=self.func, frage_text="Unterfrage?"
        )
        form = Anlage2ReviewForm()  # Formular nach dem Erstellen der Unterfrage
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"x"),
        )
        res = AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            subquestion=sub,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            subquestion=sub,
            quelle="parser",
            technisch_verfuegbar=False,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            subquestion=sub,
            quelle="ki",
            technisch_verfuegbar=True,
        )

        lookup = res.get_lookup_key()
        result_map = {lookup: res}

        row = _build_row_data(
            sub.frage_text,
            lookup,
            self.func.id,
            f"sub{sub.id}_",
            form,
            {},
            {},
            {},
            {},
            result_map,
            sub_id=sub.id,
        )

        self.assertFalse(row["doc_result"]["technisch_vorhanden"])
        self.assertTrue(row["ai_result"]["technisch_vorhanden"])


@pytest.mark.usefixtures("seed_db")
class PromptTests(NoesisTestCase):
    def test_get_prompt_returns_default(self):
        self.assertEqual(get_prompt("unknown", "foo"), "foo")

    def test_get_prompt_returns_db_value(self):
        p, _ = Prompt.objects.get_or_create(
            name="classify_system", defaults={"text": "orig"}
        )
        p.text = "DB"
        p.save()
        self.assertEqual(get_prompt("classify_system", "x"), "DB")


    def test_gap_report_anlage2_placeholder(self):
        p = Prompt.objects.get(name="gap_report_anlage2")
        self.assertIn("{gap_list}", p.text)
        self.assertIn("{system_name}", p.text)
        self.assertNotIn("{funktionen}", p.text)

    # check_anlage3_vision Prompt entfernt – kein Test mehr erforderlich



class CheckAnlage5Tests(NoesisTestCase):
    def test_check_anlage5_sets_flag(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=5,
            upload=SimpleUploadedFile("a.docx", b""),
            text_content="",
        )
        purposes = list(ZweckKategorieA.objects.values_list("beschreibung", flat=True))
        text = " ".join(purposes)
        expected_ids = list(ZweckKategorieA.objects.values_list("pk", flat=True))
        with patch("core.llm_tasks.extract_text", return_value=text):
            data = check_anlage5(pf.pk)
        pf.refresh_from_db()
        review = pf.anlage5review
        self.assertEqual(set(data["purposes"]), set(expected_ids))
        self.assertTrue(pf.verhandlungsfaehig)
        self.assertEqual(
            set(review.found_purposes.values_list("pk", flat=True)), set(expected_ids)
        )

    def test_check_anlage5_detects_other_text(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=5,
            upload=SimpleUploadedFile("a.docx", b""),
            text_content="",
        )
        cat = ZweckKategorieA.objects.create(beschreibung="A")
        text = f"{cat.beschreibung} Sonstige Zwecke zur Leistungs- oder und Verhaltenskontrolle Test"
        with patch("core.llm_tasks.extract_text", return_value=text):
            data = check_anlage5(pf.pk)
        pf.refresh_from_db()
        self.assertFalse(pf.verhandlungsfaehig)
        self.assertEqual(data["sonstige"], "Test")


class PromptImportTests(NoesisTestCase):
    def setUp(self):
        admin_group = Group.objects.create(name="admin")
        self.user = User.objects.create_user("pimport", password="pass")
        self.user.groups.add(admin_group)
        self.client.login(username="pimport", password="pass")

    def test_import_with_clear_first_replaces_prompts(self):
        Prompt.objects.create(name="old", text="x")
        payload = json.dumps([{"name": "neu", "text": "t"}])
        file = SimpleUploadedFile("p.json", payload.encode("utf-8"))
        url = reverse("admin_prompt_import")
        resp = self.client.post(
            url,
            {"json_file": file, "clear_first": "on"},
            format="multipart",
        )
        self.assertRedirects(resp, reverse("admin_prompts"))
        self.assertEqual(Prompt.objects.count(), 1)
        self.assertTrue(Prompt.objects.filter(name="neu").exists())


class ReportingTests(NoesisTestCase):
    def test_gap_analysis_file_created(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("a.txt", b"data"),
            text_content="Testtext",
            analysis_json={"ok": {"value": True, "editable": True}},
        )
        path = generate_gap_analysis(projekt)
        try:
            self.assertTrue(path.exists())
        finally:
            path.unlink(missing_ok=True)


class ProjektFileCheckViewTests(NoesisTestCase):
    def setUp(self):
        self.user = User.objects.create_user("user2", password="pass")
        self.client.login(username="user2", password="pass")
        self.projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("a.txt", b"data"),
            text_content="Text",
            analysis_json={},
        )

    def test_file_check_endpoint_saves_json(self):
        url = reverse("projekt_file_check", args=[self.projekt.pk, 1])
        resp = self.client.post(url)
        self.assertEqual(resp.status_code, 200)
        resp_json = resp.json()
        file_obj = self.projekt.anlagen.get(anlage_nr=1)
        expected = {"questions": {}}
        self.assertEqual(file_obj.analysis_json, expected)
        self.assertEqual(resp_json["analysis"], expected)

    def test_file_check_pk_endpoint_saves_json(self):
        file_obj = self.projekt.anlagen.get(anlage_nr=1)
        url = reverse("projekt_file_check_pk", args=[file_obj.pk])
        resp = self.client.post(url)
        self.assertEqual(resp.status_code, 200)
        resp_json = resp.json()
        file_obj.refresh_from_db()
        expected = {"questions": {}}
        self.assertEqual(file_obj.analysis_json, expected)
        self.assertEqual(resp_json["analysis"], expected)


class Anlage2ReviewTests(NoesisTestCase):
    def setUp(self):
        super().setUp()
        patcher = patch("core.signals.start_analysis_for_file", return_value="")
        patcher.start()
        self.addCleanup(patcher.stop)
        self.user = User.objects.create_user("rev", password="pass")
        self.client.login(username="rev", password="pass")
        self.projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        self.file = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("c.txt", b"d"),
            text_content="Text",
            analysis_json={
                "functions": [
                    {
                        "funktion": "Anmelden",
                        "technisch_vorhanden": {"value": True, "note": None},
                        "einsatz_bei_telefonica": {"value": False, "note": None},
                        "zur_lv_kontrolle": {"value": False, "note": None},
                        "ki_beteiligung": {"value": True, "note": None},
                    }
                ]
            },
        )
        self.func = Anlage2Function.objects.create(name="Anmelden")
        Anlage2SubQuestion.objects.filter(funktion=self.func).delete()
        self.sub = Anlage2SubQuestion.objects.create(
            funktion=self.func, frage_text="Warum?"
        )

    def test_get_shows_table(self):
        url = reverse("projekt_file_edit_json", args=[self.file.pk])
        resp = self.client.get(url, follow=True)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Anlage 2 Funktionen prüfen")

    def test_post_saves_data(self):
        url = reverse("projekt_file_edit_json", args=[self.file.pk])
        resp = self.client.post(
            url,
            {
                f"func{self.func.id}_technisch_vorhanden": "on",
                f"sub{self.sub.id}_ki_beteiligung": "on",
            },
        )
        self.assertRedirects(resp, reverse("projekt_detail", args=[self.projekt.pk]))
        self.file.refresh_from_db()
        fe_func = FunktionsErgebnis.objects.filter(
            anlage_datei=self.file,
            funktion=self.func,
            subquestion__isnull=True,
            quelle="manuell",
            technisch_verfuegbar=True,
        ).first()
        self.assertIsNotNone(fe_func)
        self.assertTrue(fe_func.technisch_verfuegbar)
        fe_sub = FunktionsErgebnis.objects.filter(
            anlage_datei=self.file,
            funktion=self.func,
            subquestion=self.sub,
            quelle="manuell",
            ki_beteiligung=True,
        ).first()
        self.assertIsNotNone(fe_sub)
        self.assertTrue(fe_sub.ki_beteiligung)

    def test_prefill_from_analysis(self):
        """Die Formulardaten verwenden Analysewerte als Vorgabe."""
        self.file.analysis_json = {
            "functions": [
                {
                    "funktion": "Anmelden",
                    "technisch_vorhanden": {"value": True, "note": None},
                    "einsatz_bei_telefonica": {"value": True, "note": None},
                    "zur_lv_kontrolle": {"value": True, "note": None},
                }
            ]
        }
        self.file.save()

        url = reverse("projekt_file_edit_json", args=[self.file.pk])
        resp = self.client.get(url)
        field = f"func{self.func.id}_technisch_vorhanden"
        self.assertTrue(resp.context["form"].initial[field])

    def test_prefill_with_metadaten_no_ergebnis(self):
        """Analysewerte werden auch ohne FunktionsErgebnisse angezeigt."""
        # Keine manuellen Werte gesetzt – nur Metadaten vorhanden
        self.file.save()
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=self.file,
            funktion=self.func,
        )
        url = reverse("projekt_file_edit_json", args=[self.file.pk])
        resp = self.client.get(url)
        field = f"func{self.func.id}_technisch_vorhanden"
        self.assertTrue(resp.context["form"].initial[field])

    def test_rows_include_lookup_key(self):
        url = reverse("projekt_file_edit_json", args=[self.file.pk])
        resp = self.client.get(url)
        rows = resp.context["rows"]
        # Verif-Key unabhängig von der Reihenfolge prüfen
        verif_keys = [row["verif_key"] for row in rows]
        self.assertIn(self.func.name, verif_keys)
        self.assertIn(f"{self.func.name}: {self.sub.frage_text}", verif_keys)

    def test_subquestion_justification_link(self):
        FunktionsErgebnis.objects.create(
            anlage_datei=self.file,
            funktion=self.func,
            subquestion=self.sub,
            quelle="ki",
            begruendung="Text",
        )

        url = reverse("projekt_file_edit_json", args=[self.file.pk])
        resp = self.client.get(url)
        link = reverse(
            "justification_detail_edit",
            args=[self.file.pk, f"{self.func.name}: {self.sub.frage_text}"],
        )
        self.assertContains(resp, link)

    def test_no_auto_analysis_on_get(self):
        pf = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("n.txt", b"x"),
            text_content="t",
        )
        url = reverse("projekt_file_edit_json", args=[pf.pk])

        def _fake(obj):
            obj.analysis_json = {"functions": []}
            obj.save(update_fields=["analysis_json"])
            return []

        with patch("core.views.run_anlage2_analysis", side_effect=_fake) as mock_func:
            self.client.get(url)
            self.client.get(url)
        self.assertEqual(mock_func.call_count, 0)


class WorkerGenerateGutachtenTests(NoesisTestCase):
    def setUp(self):
        self.projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("a.txt", b"data"),
            text_content="Text",
            analysis_json={},
        )
        self.knowledge = SoftwareKnowledge.objects.create(
            project=self.projekt,
            software_name="A",
            is_known_by_llm=True,
            description="",
        )

    def test_worker_creates_file(self):
        with patch("core.llm_tasks.query_llm", return_value="Text"):
            path = worker_generate_gutachten(self.projekt.pk, self.knowledge.pk)
        self.projekt.refresh_from_db()
        self.assertTrue(self.projekt.gutachten_file.name)
        self.assertEqual(
            Gutachten.objects.filter(software_knowledge=self.knowledge).count(), 1
        )
        Path(path).unlink(missing_ok=True)

    def test_worker_updates_existing_gutachten(self):
        Gutachten.objects.create(software_knowledge=self.knowledge, text="Alt")
        with patch("core.llm_tasks.query_llm", return_value="Neu"):
            path = worker_generate_gutachten(self.projekt.pk, self.knowledge.pk)
        gutachten = Gutachten.objects.get(software_knowledge=self.knowledge)
        self.assertEqual(gutachten.text, "Neu")
        self.assertEqual(
            Gutachten.objects.filter(software_knowledge=self.knowledge).count(), 1
        )
        Path(path).unlink(missing_ok=True)


class ProjektFileDeleteResultTests(NoesisTestCase):
    def setUp(self):
        self.user = User.objects.create_user("deluser", password="pass")
        self.client.login(username="deluser", password="pass")
        self.projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        self.file = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=3,
            upload=SimpleUploadedFile("d.txt", b"data"),
            text_content="Text",
            analysis_json={
                "verhandlungsfaehig": {"value": True},
                "pages": {"value": 1},
            },
            manual_reviewed=True,
            verhandlungsfaehig=True,
        )

    def test_delete_resets_fields(self):
        url = reverse("projekt_file_delete_result", args=[self.file.pk])
        resp = self.client.post(url)
        self.assertRedirects(resp, reverse("anlage3_review", args=[self.projekt.pk]))
        self.file.refresh_from_db()
        self.assertIsNone(self.file.analysis_json)
        self.assertFalse(self.file.manual_reviewed)
        self.assertFalse(self.file.verhandlungsfaehig)


@pytest.mark.usefixtures("seed_db")
class ProjektFileVersionDeletionTests(NoesisTestCase):
    def setUp(self):
        self.projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        self.superuser = User.objects.get(username="frank")

    def test_delete_active_restores_parent(self):
        self.client.login(username=self.superuser.username, password="pass")
        v1 = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("a.txt", b"d"),
            text_content="t",
        )
        v1.is_active = False
        v1.save(update_fields=["is_active"])
        v2 = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("b.txt", b"d"),
            text_content="t",
            version=2,
            parent=v1,
        )
        url = reverse("delete_project_file_version", args=[v2.pk])
        resp = self.client.post(url, follow=True)
        self.assertRedirects(resp, reverse("projekt_detail", args=[self.projekt.pk]))
        self.assertContains(resp, "Die Version wurde erfolgreich gel\u00f6scht.")
        self.assertFalse(BVProjectFile.objects.filter(pk=v2.pk).exists())
        v1.refresh_from_db()
        self.assertTrue(v1.is_active)

    def test_delete_inactive_repairs_chain(self):
        self.client.login(username=self.superuser.username, password="pass")
        v1 = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("a.txt", b"d"),
            text_content="t",
            is_active=False,
        )
        v2 = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("b.txt", b"d"),
            text_content="t",
            version=2,
            parent=v1,
            is_active=False,
        )
        v3 = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("c.txt", b"d"),
            text_content="t",
            version=3,
            parent=v2,
        )
        url = reverse("delete_project_file_version", args=[v2.pk])
        resp = self.client.post(url, follow=True)
        self.assertContains(resp, "Die Version wurde erfolgreich gel\u00f6scht.")
        self.assertFalse(BVProjectFile.objects.filter(pk=v2.pk).exists())
        v3.refresh_from_db()
        self.assertEqual(v3.parent, v1)


class ProjektFileCheckResultTests(NoesisTestCase):
    def setUp(self):
        self.user = User.objects.create_user("vuser", password="pass")
        self.client.login(username="vuser", password="pass")
        self.projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        self.file = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("a.txt", b"data"),
            text_content="Text",
            analysis_json={},
        )
        self.file2 = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("b.txt", b"data"),
            text_content="Text2",
            analysis_json={},
        )

    def test_get_runs_check_and_redirects_to_edit(self):
        url = reverse("projekt_file_check_view", args=[self.file.pk])
        resp = self.client.get(url)
        self.assertRedirects(
            resp, reverse("projekt_file_edit_json", args=[self.file.pk])
        )
        self.file.refresh_from_db()
        expected = {"questions": {}}
        self.assertEqual(self.file.analysis_json, expected)

    def test_post_triggers_check_and_redirects(self):
        url = reverse("projekt_file_check_view", args=[self.file.pk])
        with patch("core.views.check_anlage1") as mock_func:
            mock_func.return_value = {"questions": {}}
            resp = self.client.post(url)
        self.assertRedirects(
            resp, reverse("projekt_file_edit_json", args=[self.file.pk])
        )
        mock_func.assert_called_with(self.file.pk)

    def test_anlage2_uses_parser_by_default(self):
        url = reverse("projekt_file_check_view", args=[self.file2.pk])
        with patch("core.views.run_anlage2_analysis") as mock_func:
            mock_func.return_value = []
            resp = self.client.get(url)
        self.assertRedirects(
            resp, reverse("projekt_file_edit_json", args=[self.file2.pk])
        )
        mock_func.assert_called_with(self.file2)

    def test_llm_param_triggers_full_check(self):
        url = reverse("projekt_file_check_view", args=[self.file2.pk]) + "?llm=1"
        with patch("core.views.check_anlage2") as mock_func:
            mock_func.return_value = {"task": "check_anlage2"}
            resp = self.client.get(url)
        self.assertRedirects(
            resp, reverse("projekt_file_edit_json", args=[self.file2.pk])
        )
        mock_func.assert_called_with(self.file2.pk)

    def test_anlage3_uses_analysis(self):
        pf = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=3,
            upload=SimpleUploadedFile("c.txt", b"x"),
            text_content="T",
        )
        url = reverse("projekt_file_check_view", args=[pf.pk])
        with patch("core.views.analyse_anlage3") as mock_func:
            mock_func.return_value = {"task": "analyse_anlage3"}
            resp = self.client.get(url)
        self.assertRedirects(resp, reverse("anlage3_review", args=[self.projekt.pk]))
        mock_func.assert_called_with(pf.pk)

    def test_anlage3_llm_param_calls_analyse(self):
        pf = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=3,
            upload=SimpleUploadedFile("d.txt", b"x"),
            text_content="T",
        )
        url = reverse("projekt_file_check_view", args=[pf.pk]) + "?llm=1"
        with patch("core.views.analyse_anlage3") as mock_func:
            mock_func.return_value = {"task": "analyse_anlage3"}
            resp = self.client.get(url)
        self.assertRedirects(resp, reverse("anlage3_review", args=[self.projekt.pk]))
        mock_func.assert_called_with(pf.pk)

    def test_parse_view_runs_parser(self):
        url = reverse("projekt_file_check_view", args=[self.file2.pk])
        with patch("core.views.run_anlage2_analysis") as mock_func:
            mock_func.return_value = []
            resp = self.client.get(url)
        self.assertRedirects(
            resp, reverse("projekt_file_edit_json", args=[self.file2.pk])
        )
        mock_func.assert_called_with(self.file2)

    def test_parse_view_rejects_other_files(self):
        url = reverse("projekt_file_parse_anlage2", args=[self.file.pk])
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 404)


class LLMConfigTests(NoesisTestCase):
    @override_settings(GOOGLE_API_KEY="x")
    @patch("google.generativeai.list_models")
    @patch("google.generativeai.configure")
    def test_ready_populates_models(self, mock_conf, mock_list):
        mock_list.return_value = [
            type("M", (), {"name": "m1"})(),
            type("M", (), {"name": "m2"})(),
        ]
        from core.signals import init_llm_config

        init_llm_config(None)
        cfg = LLMConfig.objects.first()
        self.assertIsNotNone(cfg)
        self.assertEqual(cfg.available_models, ["m1", "m2"])
        self.assertTrue(cfg.models_changed)

    @override_settings(GOOGLE_API_KEY="x")
    @patch("google.generativeai.list_models")
    @patch("google.generativeai.configure")
    def test_ready_updates_models(self, mock_conf, mock_list):
        mock_list.return_value = [type("M", (), {"name": "new"})()]
        from core.signals import init_llm_config

        cfg = LLMConfig.objects.first()
        cfg.available_models = ["old"]
        cfg.models_changed = False
        cfg.save()

        init_llm_config(None)
        cfg.refresh_from_db()
        self.assertEqual(cfg.available_models, ["new"])
        self.assertTrue(cfg.models_changed)


class Anlage2ConfigSingletonTests(NoesisTestCase):
    def test_single_instance_enforced(self):
        first = Anlage2Config.get_instance()
        from django.db import transaction

        with self.assertRaises(IntegrityError):
            with transaction.atomic():
                Anlage2Config.objects.create()
        self.assertEqual(Anlage2Config.objects.count(), 1)


class TileVisibilityTests(NoesisTestCase):
    def setUp(self):
        admin_group = Group.objects.create(name="admin")
        self.user = User.objects.create_user("tileuser", password="pass")
        self.user.groups.add(admin_group)
        self.group = Group.objects.create(name="role")
        work = Area.objects.get_or_create(slug="work", defaults={"name": "Work"})[0]
        self.personal = Area.objects.get_or_create(
            slug="personal", defaults={"name": "Personal"}
        )[0]
        self.talkdiary = Tile.objects.get_or_create(
            slug="talkdiary",
            defaults={
                "name": "TalkDiary",
                "url_name": "talkdiary_personal",
            },
        )[0]
        self.talkdiary.areas.add(self.personal)
        self.group.areas.add(self.personal)
        self.group.tiles.add(self.talkdiary)
        self.projekt = Tile.objects.get_or_create(
            slug="projektverwaltung",
            defaults={
                "name": "Projektverwaltung",
                "url_name": "projekt_list",
            },
        )[0]
        self.projekt.areas.add(work)
        self.group.areas.add(work)
        self.cfg = LLMConfig.objects.first() or LLMConfig.objects.create(
            models_changed=False
        )
        self.client.login(username="tileuser", password="pass")

    def test_personal_without_access(self):
        resp = self.client.get(reverse("personal"))
        self.assertNotContains(resp, "TalkDiary")

    def test_personal_with_access(self):
        UserTileAccess.objects.create(user=self.user, tile=self.talkdiary)
        resp = self.client.get(reverse("personal"))
        self.assertContains(resp, "TalkDiary")

    def test_personal_with_group_access(self):
        self.user.groups.add(self.group)
        resp = self.client.get(reverse("personal"))
        self.assertContains(resp, "TalkDiary")

    def test_personal_with_image(self):
        UserTileAccess.objects.create(user=self.user, tile=self.talkdiary)
        self.talkdiary.image.save(
            "img.png",
            SimpleUploadedFile("img.png", b"data"),
            save=True,
        )
        resp = self.client.get(reverse("personal"))
        self.assertContains(resp, "<img", html=False)

    def _login(self, name: str) -> User:
        """Erzeugt einen Benutzer und loggt ihn ein."""
        user = User.objects.create_user(name, password="pass")
        self.client.login(username=name, password="pass")
        return user

    def test_talkdiary_access_denied_without_tile(self):
        self._login("nodiac")
        resp = self.client.get(reverse("personal"))
        self.assertNotContains(resp, "TalkDiary")
        resp = self.client.get(reverse("talkdiary_personal"))
        self.assertEqual(resp.status_code, 403)

    def test_talkdiary_access_allowed_with_tile(self):
        user = self._login("withdia")
        UserTileAccess.objects.create(user=user, tile=self.talkdiary)
        resp = self.client.get(reverse("personal"))
        self.assertContains(resp, "TalkDiary")
        resp = self.client.get(reverse("talkdiary_personal"))
        self.assertEqual(resp.status_code, 200)

    def test_projekt_tile_hidden_without_group(self):
        self._login("noproj")
        resp = self.client.get(reverse("work"))
        self.assertNotContains(resp, "Projektverwaltung")
        resp = self.client.get(reverse("projekt_list"))
        self.assertEqual(resp.status_code, 200)

    def test_projekt_access_allowed_with_tile(self):
        user = self._login("withproj")
        UserTileAccess.objects.create(user=user, tile=self.projekt)
        resp = self.client.get(reverse("work"))
        self.assertContains(resp, "Projektverwaltung")
        resp = self.client.get(reverse("projekt_list"))
        self.assertEqual(resp.status_code, 200)

    def test_projekt_tile_visible_with_group(self):
        user = self._login("groupproj")
        user.groups.add(self.group)
        GroupTileAccess.objects.create(group=self.group, tile=self.projekt)
        resp = self.client.get(reverse("work"))
        self.assertContains(resp, "Projektverwaltung")
        resp = self.client.get(reverse("projekt_list"))
        self.assertEqual(resp.status_code, 200)


class LLMConfigNoticeMiddlewareTests(NoesisTestCase):
    def setUp(self):
        admin_group = Group.objects.create(name="admin")
        # Middleware benötigt einen Staff-Benutzer
        self.user = User.objects.create_user(
            "llmadmin", password="pass", is_staff=True
        )
        self.user.groups.add(admin_group)
        self.client.login(username="llmadmin", password="pass")
        cfg, _ = LLMConfig.objects.get_or_create()
        cfg.models_changed = True
        cfg.save()

    def test_message_shown(self):
        resp = self.client.get(reverse("home"))
        msgs = [m.message for m in resp.context["messages"]]
        self.assertTrue(any("LLM-Einstellungen" in m for m in msgs))


class HomeRedirectTests(NoesisTestCase):
    def setUp(self):
        self.user = User.objects.create_user("redir", password="pass")
        personal = Area.objects.get_or_create(
            slug="personal", defaults={"name": "Personal"}
        )[0]
        tile = Tile.objects.get_or_create(
            slug="talkdiary",
            defaults={
                "name": "TalkDiary",
                "url_name": "talkdiary_personal",
            },
        )[0]
        tile.areas.add(personal)
        UserTileAccess.objects.create(user=self.user, tile=tile)
        self.client.login(username="redir", password="pass")

    def test_redirect_personal(self):
        resp = self.client.get(reverse("home"))
        self.assertRedirects(resp, reverse("personal"))


class AreaImageTests(NoesisTestCase):
    def setUp(self):
        self.user = User.objects.create_user("areauser", password="pass")
        self.client.login(username="areauser", password="pass")

    def test_home_without_images(self):
        Area.objects.get_or_create(slug="work", defaults={"name": "Work"})
        Area.objects.get_or_create(slug="personal", defaults={"name": "Personal"})
        resp = self.client.get(reverse("home"))
        self.assertNotContains(resp, 'alt="Work"', html=False)
        self.assertNotContains(resp, 'alt="Personal"', html=False)

    def test_home_with_images(self):
        work, _ = Area.objects.get_or_create(slug="work", defaults={"name": "Work"})
        personal, _ = Area.objects.get_or_create(
            slug="personal", defaults={"name": "Personal"}
        )
        work.image.save("w.png", SimpleUploadedFile("w.png", b"d"), save=True)
        personal.image.save("p.png", SimpleUploadedFile("p.png", b"d"), save=True)
        resp = self.client.get(reverse("home"))
        self.assertContains(resp, reverse("work"))
        self.assertContains(resp, reverse("personal"))
        self.assertContains(resp, "Arbeitsassistent")
        self.assertContains(resp, "Persönlicher Bereich")


class RecordingDeleteTests(NoesisTestCase):
    def setUp(self):
        self.user = User.objects.create_user("recuser", password="pass")
        self.client.login(username="recuser", password="pass")
        self.personal = Area.objects.get_or_create(
            slug="personal", defaults={"name": "Personal"}
        )[0]
        self.tile = Tile.objects.get_or_create(
            slug="talkdiary",
            defaults={
                "name": "TalkDiary",
                "url_name": "talkdiary_personal",
            },
        )[0]
        self.tile.areas.add(self.personal)
        UserTileAccess.objects.create(user=self.user, tile=self.tile)
        audio = SimpleUploadedFile("a.wav", b"data")
        transcript = SimpleUploadedFile("a.md", b"text")
        self.rec = Recording.objects.create(
            user=self.user,
            bereich=self.personal,
            audio_file=audio,
            transcript_file=transcript,
        )
        self.audio_path = Path(self.rec.audio_file.path)
        self.trans_path = Path(self.rec.transcript_file.path)

    def test_delete_own_recording(self):
        url = reverse("recording_delete", args=[self.rec.pk])
        resp = self.client.post(url)
        self.assertRedirects(resp, reverse("talkdiary_personal"))
        self.assertFalse(Recording.objects.filter(pk=self.rec.pk).exists())
        self.assertFalse(self.audio_path.exists())
        self.assertFalse(self.trans_path.exists())

    def test_delete_requires_post(self):
        url = reverse("recording_delete", args=[self.rec.pk])
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 405)
        self.assertTrue(Recording.objects.filter(pk=self.rec.pk).exists())

    def test_delete_other_user_recording(self):
        other = User.objects.create_user("other", password="pass")
        rec = Recording.objects.create(
            user=other,
            bereich=self.personal,
            audio_file=SimpleUploadedFile("b.wav", b"d"),
            transcript_file=SimpleUploadedFile("b.md", b"t"),
        )
        url = reverse("recording_delete", args=[rec.pk])
        resp = self.client.post(url)
        self.assertEqual(resp.status_code, 404)
        self.assertTrue(Recording.objects.filter(pk=rec.pk).exists())

class FunctionImportExportTests(NoesisTestCase):
    def setUp(self):
        super().setUp()
        admin_group = Group.objects.create(name="admin")
        self.user = User.objects.create_user("adminie", password="pass")
        self.user.groups.add(admin_group)
        self.client.login(username="adminie", password="pass")
        self.func = Anlage2Function.objects.create(name="FuncRoundtrip")

    def test_export_returns_json(self):
        Anlage2SubQuestion.objects.create(funktion=self.func, frage_text="Warum?")
        url = reverse("anlage2_function_export")
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        data = json.loads(resp.content)
        entry = next(item for item in data if item["name"] == "FuncRoundtrip")
        self.assertEqual(entry["subquestions"][0]["frage_text"], "Warum?")

    def test_import_creates_functions(self):
        payload = json.dumps([{"name": "FuncRoundtrip", "subquestions": ["Frage"]}])
        file = SimpleUploadedFile("func.json", payload.encode("utf-8"))
        url = reverse("anlage2_function_import")
        resp = self.client.post(
            url,
            {"json_file": file, "clear_first": "on"},
            format="multipart",
        )
        self.assertRedirects(resp, reverse("anlage2_function_list"))
        self.assertTrue(Anlage2Function.objects.filter(name="FuncRoundtrip").exists())

    def test_import_accepts_german_keys(self):
        payload = json.dumps(
            [
                {
                    "funktion": "Anwesenheit",
                    "unterfragen": [{"frage": "Testfrage"}],
                }
            ]
        )
        file = SimpleUploadedFile("func_de.json", payload.encode("utf-8"))
        url = reverse("anlage2_function_import")
        resp = self.client.post(
            url,
            {"json_file": file, "clear_first": "on"},
            format="multipart",
        )
        self.assertRedirects(resp, reverse("anlage2_function_list"))
        self.assertTrue(Anlage2Function.objects.filter(name="Anwesenheit").exists())
        self.assertEqual(
            Anlage2SubQuestion.objects.filter(funktion__name="Anwesenheit").count(),
            1,
        )

    def test_roundtrip_preserves_aliases(self):
        func = self.func
        func.detection_phrases = {"name_aliases": ["Sign in"]}
        func.save()
        Anlage2SubQuestion.objects.filter(funktion=func).delete()
        Anlage2SubQuestion.objects.create(
            funktion=func,
            frage_text="Warum?",
            detection_phrases={"name_aliases": ["Why"]},
        )
        export_url = reverse("anlage2_function_export")
        resp = self.client.get(export_url)
        payload = resp.content
        file = SimpleUploadedFile("func.json", payload)
        import_url = reverse("anlage2_function_import")
        resp = self.client.post(
            import_url,
            {"json_file": file, "clear_first": "on"},
            format="multipart",
        )
        self.assertRedirects(resp, reverse("anlage2_function_list"))
        func = Anlage2Function.objects.get(name=self.func.name)
        self.assertEqual(func.detection_phrases.get("name_aliases"), ["Sign in"])
        sub = func.anlage2subquestion_set.get(frage_text="Warum?")
        self.assertEqual(sub.detection_phrases.get("name_aliases"), ["Why"])


class GutachtenLLMCheckTests(NoesisTestCase):
    def setUp(self):
        self.user = User.objects.create_user("gcheck", password="pass")
        self.client.login(username="gcheck", password="pass")
        self.projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        self.knowledge = SoftwareKnowledge.objects.create(
            project=self.projekt,
            software_name="A",
            is_known_by_llm=True,
        )
        self.gutachten = Gutachten.objects.create(
            software_knowledge=self.knowledge, text="Test"
        )

    def test_endpoint_disabled_does_not_update_note(self):
        url = reverse("gutachten_llm_check", args=[self.gutachten.pk])
        resp = self.client.post(url)
        self.assertRedirects(
            resp, reverse("projekt_initial_pruefung", args=[self.projekt.pk])
        )
        self.projekt.refresh_from_db()
        self.assertEqual(self.projekt.gutachten_function_note, "")


class FeatureVerificationTests(NoesisTestCase):
    def setUp(self):
        super().setUp()
        self.projekt = BVProject.objects.create(
            software_typen="Word, Excel",
            beschreibung="x",
        )
        self.pf = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"data"),
            analysis_json={},
        )
        self.func = Anlage2Function.objects.create(name="Export")
        Anlage2SubQuestion.objects.filter(funktion=self.func).delete()
        self.sub = Anlage2SubQuestion.objects.create(
            funktion=self.func,
            frage_text="Warum?",
        )

    def test_any_yes_returns_true(self):
        with patch(
            "core.llm_tasks.query_llm",
            side_effect=["Ja", "Nein", "Begruendung", "Nein"],
        ) as mock_q:
            result = worker_verify_feature(self.pf.pk, "function", self.func.pk)
        self.assertEqual(
            result,
            {
                "technisch_verfuegbar": True,
                "ki_begruendung": "Begruendung",
                "ki_beteiligt": False,
                "ki_beteiligt_begruendung": "",
            },
        )
        self.assertEqual(mock_q.call_count, 4)
        pf = BVProjectFile.objects.get(project=self.projekt, anlage_nr=2)
        res = AnlagenFunktionsMetadaten.objects.get(
            anlage_datei=pf,
            funktion=self.func,
        )
        fe = FunktionsErgebnis.objects.filter(
            anlage_datei=pf, funktion=self.func, quelle="ki"
        ).first()
        self.assertIsNotNone(fe)
        self.assertTrue(fe.technisch_verfuegbar)
        self.assertFalse(fe.ki_beteiligung)
        self.assertEqual(fe.begruendung, "Begruendung")

    def test_all_no_returns_false(self):
        with patch(
            "core.llm_tasks.query_llm",
            side_effect=["Nein", "Nein"],
        ):
            result = worker_verify_feature(self.pf.pk, "subquestion", self.sub.pk)
        self.assertEqual(
            result,
            {
                "technisch_verfuegbar": False,
                "ki_begruendung": "",
                "ki_beteiligt": None,
                "ki_beteiligt_begruendung": "",
            },
        )

        pf = BVProjectFile.objects.get(project=self.projekt, anlage_nr=2)
        res = AnlagenFunktionsMetadaten.objects.get(
            anlage_datei=pf,
            funktion=self.func,
        )
        fe = FunktionsErgebnis.objects.filter(
            anlage_datei=pf, funktion=self.func, subquestion=self.sub, quelle="ki"
        ).first()
        self.assertIsNotNone(fe)
        self.assertFalse(fe.technisch_verfuegbar)

    def test_subquestion_context_contains_question(self):
        """Die Subquestion wird korrekt im Kontext übergeben."""
        with patch(
            "core.llm_tasks.query_llm",
            side_effect=["Nein", "Nein"],
        ) as mock_q:
            worker_verify_feature(self.pf.pk, "subquestion", self.sub.pk)
        first_call_ctx = mock_q.call_args_list[0].args[1]
        self.assertEqual(first_call_ctx["subquestion_text"], self.sub.frage_text)

    def test_gutachten_text_is_added_to_context(self):
        doc = Document()
        doc.add_paragraph("Info")
        tmp = NamedTemporaryFile(delete=False, suffix=".docx")
        doc.save(tmp.name)
        tmp.close()
        dest_dir = Path(settings.MEDIA_ROOT) / "gutachten"
        dest_dir.mkdir(parents=True, exist_ok=True)
        dest = dest_dir / Path(tmp.name).name
        shutil.copy(tmp.name, dest)
        Path(tmp.name).unlink(missing_ok=True)
        self.projekt.gutachten_file.name = f"gutachten/{dest.name}"
        self.projekt.save(update_fields=["gutachten_file"])
        with patch(
            "core.llm_tasks.query_llm",
            side_effect=["Ja", "Nein", "B", "Nein"],
        ) as mock_q:
            worker_verify_feature(self.pf.pk, "function", self.func.pk)
        ctx = mock_q.call_args_list[0].args[1]
        self.assertIn("Info", ctx["gutachten"])
        dest.unlink(missing_ok=True)

    def test_mixed_returns_none(self):
        with patch(
            "core.llm_tasks.query_llm",
            side_effect=["Unsicher", "Nein", "Begruendung"],
        ):
            result = worker_verify_feature(self.pf.pk, "function", self.func.pk)
        self.assertIsNone(result["technisch_verfuegbar"])
        self.assertEqual(result["ki_begruendung"], "Begruendung")
        self.assertIsNone(result["ki_beteiligt"])
        self.assertEqual(result["ki_beteiligt_begruendung"], "")
        pf = BVProjectFile.objects.get(project=self.projekt, anlage_nr=2)
        fe = FunktionsErgebnis.objects.filter(
            anlage_datei=pf, funktion=self.func, quelle="ki"
        ).first()
        self.assertIsNotNone(fe)
        self.assertEqual(fe.begruendung, "Begruendung")

    def test_negotiable_set_on_match(self):
        pf = BVProjectFile.objects.get(project=self.projekt, anlage_nr=2)
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf,
            funktion=self.func,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            quelle="parser",
            technisch_verfuegbar=True,
        )
        with patch(
            "core.llm_tasks.query_llm",
            side_effect=["Ja", "Nein", "", "Nein"],
        ):
            worker_verify_feature(self.pf.pk, "function", self.func.pk)
        parser_fe = FunktionsErgebnis.objects.filter(
            anlage_datei=pf,
            funktion=self.func,
            quelle="parser",
        ).first()
        ai_fe = FunktionsErgebnis.objects.filter(
            anlage_datei=pf,
            funktion=self.func,
            quelle="ki",
        ).first()
        self.assertTrue(parser_fe.technisch_verfuegbar)
        self.assertTrue(ai_fe.technisch_verfuegbar)

    def test_negotiable_not_set_on_mismatch(self):
        pf = BVProjectFile.objects.get(project=self.projekt, anlage_nr=2)
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf,
            funktion=self.func,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            quelle="parser",
            technisch_verfuegbar=False,
        )
        with patch(
            "core.llm_tasks.query_llm",
            side_effect=["Ja", "Nein", "", "Nein"],
        ):
            worker_verify_feature(self.pf.pk, "function", self.func.pk)
        parser_fe = FunktionsErgebnis.objects.filter(
            anlage_datei=pf,
            funktion=self.func,
            quelle="parser",
        ).first()
        ai_fe = FunktionsErgebnis.objects.filter(
            anlage_datei=pf,
            funktion=self.func,
            quelle="ki",
        ).first()
        self.assertFalse(parser_fe.technisch_verfuegbar)
        self.assertTrue(ai_fe.technisch_verfuegbar)

    def test_warnung_bei_geloeschter_datei(self):
        """L\u00f6schen der Datei f\u00fchrt zu Warnung statt Ausnahme."""

        pf = BVProjectFile.objects.get(project=self.projekt, anlage_nr=2)

        original_update = AnlagenFunktionsMetadaten.objects.update_or_create

        def _wrapped(*args, **kwargs):
            obj, created = original_update(*args, **kwargs)
            pf.delete()
            return obj, created

        class _DummyQS:
            def exists(self) -> bool:  # noqa: D401 - einfache Hilfsklasse
                return False


        with patch("core.llm_tasks.query_llm", side_effect=["Nein", "Nein"]):
            with patch(
                "core.llm_tasks.BVProjectFile.objects.filter",
                return_value=_DummyQS(),
            ):
                with self.assertLogs("core.llm_tasks", level="WARNING") as cm:
                    result = worker_verify_feature(
                        self.pf.pk, "function", self.func.pk
                    )

        self.assertEqual(result, {})
        self.assertTrue(
            any("Ergebnis wird verworfen" in msg for msg in cm.output)
        )

    def test_integrity_error_logs_and_returns_empty(self):
        """Allgemeiner IntegrityError führt zu Fehler-Log und leerem Ergebnis."""
        def _raise(*args, **kwargs):
            raise IntegrityError("boom")

        with patch("core.llm_tasks.query_llm", side_effect=["Nein", "Nein"]):
            with patch(
                "core.llm_tasks.AnlagenFunktionsMetadaten.objects.update_or_create",
                side_effect=_raise,
            ):
                with self.assertLogs("core.llm_tasks", level="ERROR") as cm:
                    result = worker_verify_feature(
                        self.pf.pk, "function", self.func.pk
                    )

        self.assertEqual(result, {})
        self.assertTrue(any("Integrit" in msg for msg in cm.output))


@pytest.mark.usefixtures("seed_db")
class InitialCheckTests(NoesisTestCase):
    def setUp(self):
        self.projekt = BVProject.objects.create(software_typen="A", beschreibung="x")

    def test_known_software_stores_description(self):
        with patch(
            "core.llm_tasks.query_llm",
            side_effect=["Ja", "Beschreibung"],
        ) as mock_q:
            sk = SoftwareKnowledge.objects.create(
                project=self.projekt,
                software_name="A",
            )
            result = worker_run_initial_check(sk.pk)
        self.assertTrue(result["is_known_by_llm"])
        self.assertEqual(result["description"], "Beschreibung")
        self.assertEqual(mock_q.call_count, 2)
        sk.refresh_from_db()
        self.assertTrue(sk.is_known_by_llm)
        self.assertEqual(sk.description, "Beschreibung")

    def test_unknown_sets_flags(self):
        with patch("core.llm_tasks.query_llm", return_value="Nein") as mock_q:
            sk = SoftwareKnowledge.objects.create(
                project=self.projekt,
                software_name="A",
            )
            result = worker_run_initial_check(sk.pk)
        self.assertFalse(result["is_known_by_llm"])
        self.assertEqual(result["description"], "")
        self.assertEqual(mock_q.call_count, 1)
        sk.refresh_from_db()
        self.assertFalse(sk.is_known_by_llm)
        self.assertEqual(sk.description, "")

    def test_context_is_passed_to_prompt(self):
        with patch("core.llm_tasks.query_llm", return_value="Nein") as mock_q:
            sk = SoftwareKnowledge.objects.create(
                project=self.projekt,
                software_name="A",
            )
            worker_run_initial_check(sk.pk, user_context="Hint")
        context_data = mock_q.call_args[0][1]
        assert context_data["user_context"] == "Hint"
        # Prompt mit Kontext-Variante ist deaktiviert; Standard-Prompt wird genutzt
        assert mock_q.call_args[0][0].name == "initial_check_knowledge"


class EditKIJustificationTests(NoesisTestCase):
    def setUp(self):
        super().setUp()
        self.user = User.objects.create_user("justi", password="pass")
        self.client.login(username="justi", password="pass")
        self.projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        self.file = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"data"),
            analysis_json={},
        )
        self.func = Anlage2Function.objects.create(name="Export")
        FunktionsErgebnis.objects.create(
            anlage_datei=self.file,
            funktion=self.func,
            quelle="ki",
            begruendung="Alt",
        )

    def test_get_returns_form(self):
        url = (
            reverse(
                "edit_ki_justification",
                args=[self.file.pk],
            )
            + f"?function={self.func.pk}"
        )
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Alt")

    def test_post_updates_value(self):
        url = reverse("edit_ki_justification", args=[self.file.pk])
        resp = self.client.post(
            url,
            {"function": self.func.pk, "ki_begruendung": "Neu"},
        )
        self.assertRedirects(
            resp, reverse("projekt_file_edit_json", args=[self.file.pk])
        )
        fe = FunktionsErgebnis.objects.get(
            anlage_datei=self.file, funktion=self.func, quelle="ki"
        )
        self.assertEqual(fe.begruendung, "Neu")


class JustificationDetailEditTests(NoesisTestCase):
    def setUp(self):
        super().setUp()
        self.user = User.objects.create_user("jdet", password="pw")
        self.client.login(username="jdet", password="pw")
        self.projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        self.file = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("jd.txt", b"data"),
            analysis_json={},
        )
        self.func = Anlage2Function.objects.create(name="Export")
        FunktionsErgebnis.objects.create(
            anlage_datei=self.file,
            funktion=self.func,
            quelle="ki",
            begruendung="Begruendung",
        )

    def test_get_loads_text(self):
        url = reverse("justification_detail_edit", args=[self.file.pk, self.func.name])
        resp = self.client.get(url)
        self.assertContains(resp, "Begruendung")


class KIInvolvementDetailEditTests(NoesisTestCase):
    def setUp(self):
        super().setUp()
        self.user = User.objects.create_user("kid", password="pw")
        self.client.login(username="kid", password="pw")
        self.projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        self.file = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("kid.txt", b"data"),
            analysis_json={},
        )
        self.func = Anlage2Function.objects.create(name="Export")
        FunktionsErgebnis.objects.create(
            anlage_datei=self.file,
            funktion=self.func,
            quelle="ki",
            ki_beteiligt_begruendung="Initial",
        )

    def test_get_loads_text(self):
        url = reverse("ki_involvement_detail_edit", args=[self.file.pk, self.func.name])
        resp = self.client.get(url)
        self.assertContains(resp, "Initial")

    def test_post_updates_value(self):
        url = reverse("ki_involvement_detail_edit", args=[self.file.pk, self.func.name])
        resp = self.client.post(url, {"justification": "Neu"})
        self.assertRedirects(resp, reverse("projekt_file_edit_json", args=[self.file.pk]))
        fe = FunktionsErgebnis.objects.get(
            anlage_datei=self.file, funktion=self.func, quelle="ki"
        )
        self.assertEqual(fe.ki_beteiligt_begruendung, "Neu")


class VerificationToInitialTests(NoesisTestCase):
    def setUp(self):
        super().setUp()
        self.project = BVProject.objects.create(software_typen="A", beschreibung="x")
        BVProjectFile.objects.create(
            project=self.project,
            anlage_nr=2,
            upload=SimpleUploadedFile("v.txt", b"data"),
            analysis_json={},
        )
        self.func = Anlage2Function.objects.create(name="Export")
        Anlage2SubQuestion.objects.filter(funktion=self.func).delete()
        self.sub = Anlage2SubQuestion.objects.create(
            funktion=self.func,
            frage_text="Warum?",
        )

    def test_conversion_reads_ai_fields(self):
        data = {
            "Export": {
                "technisch_verfuegbar": True,
                "ki_beteiligt": True,
                "ki_beteiligt_begruendung": "Grund",
            },
            "Export: Warum?": {
                "technisch_verfuegbar": False,
                "ki_beteiligt": False,
                "ki_beteiligt_begruendung": "Nein",
            },
        }

        pf = self.project.anlagen.get(anlage_nr=2)
        AnlagenFunktionsMetadaten.objects.create(anlage_datei=pf, funktion=self.func)
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf, funktion=self.func, subquestion=self.sub
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            quelle="ki",
            technisch_verfuegbar=True,
            ki_beteiligung=True,
            begruendung="Grund",
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            subquestion=self.sub,
            quelle="ki",
            technisch_verfuegbar=False,
            ki_beteiligung=False,
            begruendung="Nein",
        )
        result = _verification_to_initial(pf)
        func_data = result["functions"][str(self.func.id)]
        self.assertTrue(func_data["technisch_vorhanden"])
        self.assertTrue(func_data["ki_beteiligt"])
        self.assertEqual(func_data["begruendung"], "Grund")

        sub_data = func_data["subquestions"][str(self.sub.id)]
        self.assertFalse(sub_data["technisch_vorhanden"])
        self.assertFalse(sub_data["ki_beteiligt"])
        self.assertEqual(sub_data["begruendung"], "Nein")


class UserImportExportTests(NoesisTestCase):
    def setUp(self):
        admin_group = Group.objects.create(name="admin")
        self.user = User.objects.create_user("uadmin", password="pass")
        self.user.groups.add(admin_group)
        self.client.login(username="uadmin", password="pass")

        self.group = Group.objects.create(name="testgroup")
        self.area = Area.objects.get_or_create(slug="work", defaults={"name": "Work"})[
            0
        ]
        self.tile = Tile.objects.create(slug="t1", name="T", url_name="tile")
        self.tile.areas.add(self.area)
        self.group.areas.add(self.area)
        self.group.tiles.add(self.tile)

    def test_export_json(self):
        self.user.groups.add(self.group)


        url = reverse("admin_export_users_permissions")
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        data = json.loads(resp.content)
        entry = next(u for u in data if u["username"] == "uadmin")
        self.assertIn("testgroup", entry["groups"])
        self.assertIn("tile", entry["tiles"])

    def test_import_creates_user(self):
        payload = json.dumps(
            [
                {
                    "username": "neu",
                    "email": "a@b.c",
                    "groups": ["testgroup"],

                }
            ]
        )
        file = SimpleUploadedFile("u.json", payload.encode("utf-8"))
        url = reverse("admin_import_users_permissions")
        resp = self.client.post(url, {"json_file": file}, format="multipart")
        self.assertRedirects(resp, reverse("admin_user_list"))
        user = User.objects.get(username="neu")
        self.assertTrue(user.groups.filter(name="testgroup").exists())



class Anlage1ImportTests(NoesisTestCase):
    def setUp(self):
        admin_group = Group.objects.create(name="admin")
        self.user = User.objects.create_user("a1import", password="pass")
        self.user.groups.add(admin_group)
        self.client.login(username="a1import", password="pass")

    def test_clear_first_resets_questions(self):
        Anlage1Question.objects.create(num=99, text="Alt?", enabled=True)
        payload = json.dumps([{"text": "Neu?"}])
        file = SimpleUploadedFile("a1.json", payload.encode("utf-8"))
        url = reverse("admin_anlage1_import")
        resp = self.client.post(
            url,
            {"json_file": file, "clear_first": "on"},
            format="multipart",
        )
        self.assertRedirects(resp, reverse("admin_anlage1"))
        self.assertEqual(Anlage1Question.objects.count(), 1)
        q = Anlage1Question.objects.first()
        self.assertEqual(q.text, "Neu?")
        self.assertEqual(q.num, 1)


class Anlage2ConfigImportExportTests(NoesisTestCase):
    def setUp(self):
        admin_group = Group.objects.create(name="admin")
        self.user = User.objects.create_user("cfgadmin", password="pass")
        self.user.groups.add(admin_group)
        self.client.login(username="cfgadmin", password="pass")
        # Anlage2Config ist ein Singleton. Für die Tests müssen wir die
        # globale Instanz verwenden und die benötigten Felder direkt darauf
        # setzen.
        self.cfg = Anlage2Config.get_instance()
        self.cfg.parser_mode = "auto"
        self.cfg.parser_order = ["exact"]
        self.cfg.text_technisch_verfuegbar_true = ["ja"]
        self.cfg.save()

    def test_export_contains_headings_and_phrases(self):
        Anlage2ColumnHeading.objects.create(
            config=self.cfg,
            field_name="technisch_vorhanden",
            text="Verfügbar?",
        )
        AntwortErkennungsRegel.objects.create(
            regel_name="R1",
            erkennungs_phrase="ja",
            actions_json=[{"field": "technisch_verfuegbar", "value": True}],
            prioritaet=0,
        )
        a4 = Anlage4ParserConfig.objects.first()
        a4.delimiter_phrase = "X"
        a4.save()
        url = reverse("admin_anlage2_config_export")
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        data = json.loads(resp.content)
        self.assertEqual(data["config"]["parser_mode"], self.cfg.parser_mode)
        self.assertEqual(data["config"]["parser_order"], self.cfg.parser_order)
        self.assertIn(
            {"field_name": "technisch_vorhanden", "text": "Verfügbar?"},
            data["alias_headings"],
        )
        self.assertIn("answer_rules", data)
        self.assertTrue(any(r["regel_name"] == "R1" for r in data["answer_rules"]))
        self.assertEqual(data["a4_parser"]["delimiter_phrase"], "X")

    def test_import_creates_headings(self):
        payload = json.dumps(
            {
                "config": {
                "parser_mode": "text_only",
                "parser_order": ["exact"],
                    "enforce_subquestion_override": True,
                    "text_technisch_verfuegbar_true": ["ja"],
                },
                "alias_headings": [{"field_name": "ki_beteiligung", "text": "KI?"}],
                "answer_rules": [
                    {
                        "regel_name": "R2",
                        "erkennungs_phrase": "nein",
                        "actions": [{"field": "technisch_verfuegbar", "value": False}],
                        "prioritaet": 1,
                    }
                ],
                "a4_parser": {"delimiter_phrase": "Y"},
            }
        )
        file = SimpleUploadedFile("cfg.json", payload.encode("utf-8"))
        url = reverse("admin_anlage2_config_import")
        resp = self.client.post(url, {"json_file": file}, format="multipart")
        self.assertRedirects(resp, reverse("anlage2_config"))
        self.assertTrue(
            Anlage2ColumnHeading.objects.filter(
                field_name="ki_beteiligung", text="KI?"
            ).exists()
        )
        self.assertTrue(AntwortErkennungsRegel.objects.filter(regel_name="R2").exists())
        a4_cfg = Anlage4ParserConfig.objects.first()
        self.assertEqual(a4_cfg.delimiter_phrase, "Y")
        self.cfg.refresh_from_db()
        self.assertEqual(self.cfg.parser_mode, "text_only")
        self.assertEqual(self.cfg.parser_order, ["exact"])
        self.assertTrue(self.cfg.enforce_subquestion_override)
        self.assertEqual(self.cfg.text_technisch_verfuegbar_true, ["ja"])


class Anlage2ConfigViewTests(NoesisTestCase):
    def setUp(self):
        admin = Group.objects.create(name="admin")
        self.user = User.objects.create_user("cfguser", password="pass")
        self.user.groups.add(admin)
        self.client.login(username="cfguser", password="pass")
        self.cfg = Anlage2Config.get_instance()

    def _build_general_data(self, **extra) -> dict:
        """Erstellt Grunddaten für das Anlage2Config-Formular."""
        data = {name: "" for name in Anlage2ConfigForm.OPTIONAL_JSON_FIELDS}
        data.update(extra)
        return data


    def test_update_parser_mode(self):
        url = reverse("anlage2_config")
        resp = self.client.post(
            url,
            self._build_general_data(
                parser_mode="text_only",
                parser_order=["exact"],
                action="save_general",
                active_tab="general",
            ),
        )
        self.assertRedirects(resp, url + "?tab=general")
        self.cfg.refresh_from_db()
        self.assertEqual(self.cfg.parser_order, ["exact"])

    def test_update_parser_order(self):
        url = reverse("anlage2_config")
        resp = self.client.post(
            url,
            self._build_general_data(
                parser_mode=self.cfg.parser_mode,
                parser_order=["exact"],
                action="save_general",
                active_tab="general",
            ),
        )
        self.assertRedirects(resp, url + "?tab=general")
        self.cfg.refresh_from_db()
        self.assertEqual(self.cfg.parser_order, ["exact"])

    def test_save_table_tab(self):
        url = reverse("anlage2_config")
        resp = self.client.post(
            url,
            {
                "new_field": "technisch_vorhanden",
                "new_text": "Verfügbar?",
                "action": "save_table",
                "active_tab": "table",
            },
        )
        self.assertRedirects(resp, url + "?tab=table")
        self.assertTrue(Anlage2ColumnHeading.objects.filter(text="Verfügbar?").exists())

    def test_multiline_phrases_saved(self):
        url = reverse("anlage2_config")
        resp = self.client.post(
            url,
            self._build_general_data(
                text_technisch_verfuegbar_true="ja\nokay\n",
                parser_mode=self.cfg.parser_mode,
                parser_order=self.cfg.parser_order,
                action="save_general",
                active_tab="general",
            ),
        )
        self.assertRedirects(resp, url + "?tab=general")
        self.cfg.refresh_from_db()
        self.assertEqual(self.cfg.text_technisch_verfuegbar_true, ["ja", "okay"])


class ParserRuleImportExportTests(NoesisTestCase):
    def setUp(self):
        admin_group = Group.objects.create(name="admin")
        self.user = User.objects.create_user(
            "ruleadmin", password="pass", is_staff=True
        )
        self.user.groups.add(admin_group)
        self.client.login(username="ruleadmin", password="pass")

    def test_export_returns_json(self):
        AntwortErkennungsRegel.objects.create(
            regel_name="R1",
            erkennungs_phrase="ja",
            actions_json=[{"field": "tech", "value": True}],
        )
        url = reverse("anlage2_parser_rule_export")
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        data = json.loads(resp.content)
        self.assertTrue(any(r["regel_name"] == "R1" for r in data))

    def test_import_creates_rule(self):
        payload = json.dumps(
            [
                {
                    "regel_name": "R2",
                    "erkennungs_phrase": "nein",
                    "actions": [{"field": "tech", "value": False}],
                }
            ]
        )
        file = SimpleUploadedFile("rules.json", payload.encode("utf-8"))
        url = reverse("anlage2_parser_rule_import")
        resp = self.client.post(url, {"json_file": file}, format="multipart")
        self.assertRedirects(resp, reverse("parser_rule_list"))
        self.assertTrue(AntwortErkennungsRegel.objects.filter(regel_name="R2").exists())


class AjaxAnlage2ReviewTests(NoesisTestCase):
    def setUp(self):
        super().setUp()
        self.user = User.objects.create_user("reviewer", password="pw", is_staff=True)
        self.client.login(username="reviewer", password="pw")
        self.projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        self.pf = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"x"),
            analysis_json={},
        )
        self.func = Anlage2Function.objects.create(name="Anmelden")
        Anlage2SubQuestion.objects.filter(funktion=self.func).delete()

    def test_manual_result_saved(self):
        url = reverse("ajax_save_anlage2_review")
        resp = self.client.post(
            url,
            data=json.dumps(
                {
                    "project_file_id": self.pf.pk,
                    "function_id": self.func.pk,
                    "status": True,
                    "field_name": "technisch_vorhanden",
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)

        # Ergebnis wurde als manueller Eintrag gespeichert

        fe = FunktionsErgebnis.objects.filter(
            anlage_datei__project=self.projekt,
            funktion=self.func,
            quelle="manuell",
        ).first()
        self.assertTrue(fe.technisch_verfuegbar)

    def test_gap_generated_on_difference(self):
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=self.pf,
            funktion=self.func,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=self.pf,
            funktion=self.func,
            quelle="ki",
            technisch_verfuegbar=True,
        )
        url = reverse("ajax_save_anlage2_review")
        resp = self.client.post(
            url,
            data=json.dumps(
                {
                    "project_file_id": self.pf.pk,
                    "function_id": self.func.pk,
                    "status": False,
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)
        data = resp.json()
        self.assertEqual(data.get("gap_summary"), "")
        res = AnlagenFunktionsMetadaten.objects.get(
            anlage_datei=self.pf,
            funktion=self.func,
        )
        self.assertEqual(res.gap_summary, "")

        gap_url = reverse("ajax_generate_gap_summary", args=[res.pk])

        def immediate(name, *args):
            self.assertEqual(name, "core.llm_tasks.worker_generate_gap_summary")
            worker_generate_gap_summary(*args)

        with patch("core.views.async_task", side_effect=immediate):
            resp = self.client.post(gap_url)
        self.assertEqual(resp.status_code, 200)
        res.refresh_from_db()
        self.assertEqual(res.gap_notiz, "")
        self.assertEqual(res.gap_summary, "")
        gap_entry = FunktionsErgebnis.objects.filter(
            anlage_datei=self.pf,
            funktion=self.func,
            quelle="gap",
        ).latest("created_at")
        self.assertIsNotNone(gap_entry)

    def test_manual_sets_negotiable(self):
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=self.pf,
            funktion=self.func,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=self.pf,
            funktion=self.func,
            quelle="ki",
            technisch_verfuegbar=True,
        )

        url = reverse("ajax_save_anlage2_review")
        resp = self.client.post(
            url,
            data=json.dumps(
                {
                    "project_file_id": self.pf.pk,
                    "function_id": self.func.pk,
                    "status": True,
                    "field_name": "technisch_vorhanden",
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)
        fe = FunktionsErgebnis.objects.filter(
            anlage_datei__project=self.projekt,
            funktion=self.func,
            quelle="manuell",
        ).first()
        self.assertTrue(fe.technisch_verfuegbar)

    def test_save_einsatz_telefonica(self):
        url = reverse("ajax_save_anlage2_review")
        resp = self.client.post(
            url,
            data=json.dumps(
                {
                    "project_file_id": self.pf.pk,
                    "function_id": self.func.pk,
                    "status": True,
                    "field_name": "einsatz_bei_telefonica",
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)
        fe = FunktionsErgebnis.objects.filter(
            anlage_datei__project=self.projekt,
            funktion=self.func,
            quelle="manuell",
        ).first()
        self.assertTrue(fe.einsatz_bei_telefonica)
        # Zustand ist in FunktionsErgebnis persistiert

    def test_save_lv_kontrolle(self):
        url = reverse("ajax_save_anlage2_review")
        resp = self.client.post(
            url,
            data=json.dumps(
                {
                    "project_file_id": self.pf.pk,
                    "function_id": self.func.pk,
                    "status": False,
                    "field_name": "zur_lv_kontrolle",
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)
        fe = FunktionsErgebnis.objects.filter(
            anlage_datei__project=self.projekt,
            funktion=self.func,
            quelle="manuell",
        ).first()
        self.assertFalse(fe.zur_lv_kontrolle)
        # Zustand ist in FunktionsErgebnis persistiert

    def test_manual_result_merge(self):
        url = reverse("ajax_save_anlage2_review")
        self.client.post(
            url,
            data=json.dumps(
                {
                    "project_file_id": self.pf.pk,
                    "function_id": self.func.pk,
                    "status": True,
                    "field_name": "technisch_vorhanden",
                }
            ),
            content_type="application/json",
        )
        self.client.post(
            url,
            data=json.dumps(
                {
                    "project_file_id": self.pf.pk,
                    "function_id": self.func.pk,
                    "status": False,
                    "field_name": "ki_beteiligung",
                }
            ),
            content_type="application/json",
        )

        # Beide manuellen Einträge wurden gespeichert
        fes = FunktionsErgebnis.objects.filter(
            anlage_datei__project=self.projekt,
            funktion=self.func,
            quelle="manuell",
        )
        self.assertEqual(fes.count(), 2)

    def test_set_negotiable_override(self):
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=self.pf,
            funktion=self.func,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=self.pf,
            funktion=self.func,
            quelle="ki",
            technisch_verfuegbar=True,
        )
        url = reverse("ajax_save_anlage2_review")
        resp = self.client.post(
            url,
            data=json.dumps(
                {
                    "project_file_id": self.pf.pk,
                    "function_id": self.func.pk,
                    "set_negotiable": True,
                }
            ),
            content_type="application/json",
        )
        self.assertEqual(resp.status_code, 200)
        res = AnlagenFunktionsMetadaten.objects.get(
            anlage_datei=self.pf,
            funktion=self.func,
        )
        self.assertTrue(res.is_negotiable_manual_override)

        self.client.post(
            url,
            data=json.dumps(
                {
                    "project_file_id": self.pf.pk,
                    "function_id": self.func.pk,
                    "set_negotiable": None,
                }
            ),
            content_type="application/json",
        )
        res.refresh_from_db()
        self.assertIsNone(res.is_negotiable_manual_override)

    def test_negotiable_does_not_set_manual_value(self):
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=self.pf,
            funktion=self.func,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=self.pf,
            funktion=self.func,
            quelle="ki",
            technisch_verfuegbar=True,
        )
        url = reverse("ajax_save_anlage2_review")
        self.client.post(
            url,
            data=json.dumps(
                {
                    "project_file_id": self.pf.pk,
                    "function_id": self.func.pk,
                    "set_negotiable": True,
                }
            ),
            content_type="application/json",
        )
        res = AnlagenFunktionsMetadaten.objects.get(
            anlage_datei=self.pf,
            funktion=self.func,
        )
        self.assertFalse(
            FunktionsErgebnis.objects.filter(
                anlage_datei__project=self.projekt,
                funktion=self.func,
                quelle="manuell",
            ).exists()
        )

    def test_manual_yes_triggers_subquestion_checks(self):
        sub = Anlage2SubQuestion.objects.create(
            funktion=self.func,
            frage_text="S?",
        )
        url = reverse("ajax_save_anlage2_review")
        with patch("core.views.async_task") as mock_task:
            resp = self.client.post(
                url,
                data=json.dumps(
                    {
                        "project_file_id": self.pf.pk,
                        "function_id": self.func.pk,
                        "status": True,
                        "field_name": "technisch_vorhanden",
                    }
                ),
                content_type="application/json",
            )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(mock_task.call_count, 2)
        mock_task.assert_any_call(
            "core.llm_tasks.worker_verify_feature",
            self.pf.pk,
            "function",
            self.func.pk,
        )
        mock_task.assert_any_call(
            "core.llm_tasks.worker_verify_feature",
            self.pf.pk,
            "subquestion",
            sub.pk,
        )


class SupervisionGapTests(NoesisTestCase):
    def setUp(self) -> None:  # pragma: no cover - setup
        super().setUp()
        self.func = Anlage2Function.objects.create(name="Anmelden")

    def test_manually_negotiable_function_excluded_from_supervision(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("f.txt", b"x"),
        )
        func = self.func
        Anlage2SubQuestion.objects.filter(funktion=func).delete()
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf,
            funktion=func,
            is_negotiable_manual_override=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            quelle="parser",
            technisch_verfuegbar=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            quelle="ki",
            technisch_verfuegbar=False,
        )
        sub = Anlage2SubQuestion.objects.create(funktion=func, frage_text="S?")
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf,
            funktion=func,
            subquestion=sub,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            subquestion=sub,
            quelle="parser",
            technisch_verfuegbar=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            subquestion=sub,
            quelle="ki",
            technisch_verfuegbar=False,
        )
        groups = _build_supervision_groups(pf)
        self.assertEqual(groups, [])

    def test_subquestion_before_function_excluded_from_supervision(self):
        """Unterfrage vor Hauptfunktion: Funktion wird dennoch übersprungen."""

        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("f.txt", b"x"),
        )
        func = self.func
        Anlage2SubQuestion.objects.filter(funktion=func).delete()
        sub = Anlage2SubQuestion.objects.create(funktion=func, frage_text="S?")

        # Unterfrage zuerst anlegen, damit sie im Default-Ordering vor der Funktion steht
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf,
            funktion=func,
            subquestion=sub,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            subquestion=sub,
            quelle="parser",
            technisch_verfuegbar=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            subquestion=sub,
            quelle="ki",
            technisch_verfuegbar=False,
        )

        # Hauptfunktion nachträglich als verhandlungsfähig markieren
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf,
            funktion=func,
            is_negotiable_manual_override=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            quelle="parser",
            technisch_verfuegbar=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            quelle="ki",
            technisch_verfuegbar=False,
        )

        groups = _build_supervision_groups(pf)
        self.assertEqual(groups, [])

    def test_ai_reason_uses_function_begruendung(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("f.txt", b"x"),
        )
        func = self.func
        res = AnlagenFunktionsMetadaten.objects.create(anlage_datei=pf, funktion=func)
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            quelle="parser",
            technisch_verfuegbar=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            quelle="ki",
            technisch_verfuegbar=False,
            begruendung="Func reason",
            ki_beteiligt_begruendung="KI involvement",
        )
        row = _build_supervision_row(res, pf)
        self.assertEqual(row["ai_reason"], "Func reason")


@pytest.mark.usefixtures("seed_db")
class Anlage2ResetTests(NoesisTestCase):
    """Tests zum Zurücksetzen von Anlage-2-Ergebnissen."""

    def setUp(self):
        super().setUp()
        self.superuser = User.objects.get(username="frank")
        self.client.login(username=self.superuser.username, password="pass")
        self.func = Anlage2Function.objects.create(name="Anmelden")

    def test_run_anlage2_analysis_resets_results(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf_old = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("old.txt", b"x"),
        )
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf_old,
            funktion=self.func,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf_old,
            funktion=self.func,
            quelle="parser",
        )
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("new.txt", b"x"),
        )
        with (
            patch(
                "core.llm_tasks.parse_anlage2_table",
                return_value=[{"funktion": "Anmelden"}],
            ),
            patch("core.text_parser.parse_anlage2_text", return_value=[]),
        ):
            run_anlage2_analysis(pf)
        results = AnlagenFunktionsMetadaten.objects.filter(
            anlage_datei__project=projekt,
            subquestion__isnull=True,
        )
        self.assertEqual(results.count(), Anlage2Function.objects.count())
        fe = FunktionsErgebnis.objects.filter(
            anlage_datei__project=projekt,
            funktion=self.func,
            quelle="parser",
        ).first()
        self.assertIsNotNone(fe)

    def test_conditional_check_resets_results(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf_old = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("old.txt", b"x"),
        )
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf_old,
            funktion=self.func,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf_old,
            funktion=self.func,
            quelle="ki",
            technisch_verfuegbar=False,
        )

        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("new.txt", b"x"),
        )

        def fake(file_id, object_type, object_id, _model=None):
            pf_latest = BVProjectFile.objects.filter(
                project=projekt, anlage_nr=2
            ).first()
            if object_type == "function":
                fid = object_id
                AnlagenFunktionsMetadaten.objects.update_or_create(
                    anlage_datei=pf_latest,
                    funktion_id=fid,
                    defaults={},
                )
                FunktionsErgebnis.objects.create(
                    anlage_datei=pf_latest,
                    funktion_id=fid,
                    quelle="ki",
                    technisch_verfuegbar=True,
                )
            else:  # subquestion
                sub = Anlage2SubQuestion.objects.get(pk=object_id)
                AnlagenFunktionsMetadaten.objects.update_or_create(
                    anlage_datei=pf_latest,
                    funktion=sub.funktion,
                    subquestion=sub,
                    defaults={},
                )
                FunktionsErgebnis.objects.create(
                    anlage_datei=pf_latest,
                    funktion=sub.funktion,
                    subquestion=sub,
                    quelle="ki",
                    technisch_verfuegbar=True,
                )
            return {}

        with (
            patch("core.llm_tasks.worker_verify_feature", side_effect=fake) as mock_verify,
            patch("core.llm_tasks.async_task") as mock_async,
            patch("core.llm_tasks.result") as mock_result,
        ):
            mock_async.side_effect = lambda name, *a, **k: (
                mock_verify(*a, **k) or "tid"
            )
            mock_result.side_effect = lambda *a, **k: None
            run_conditional_anlage2_check(pf.pk)
        results = AnlagenFunktionsMetadaten.objects.filter(
            anlage_datei__project=projekt
        )
        self.assertEqual(results.count(), Anlage2Function.objects.count())
        fe = FunktionsErgebnis.objects.filter(
            anlage_datei__project=projekt,
            funktion=self.func,
            quelle="ki",
        ).first()
        self.assertTrue(fe.technisch_verfuegbar)

    def test_conditional_check_deletes_only_current_file_metadata(self):
        """Nur Metadaten der geprüften Anlage werden entfernt."""
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"x"),
        )
        other_pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=3,
            upload=SimpleUploadedFile("b.txt", b"x"),
        )
        func = self.func
        AnlagenFunktionsMetadaten.objects.create(anlage_datei=pf, funktion=func)
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=other_pf, funktion=func
        )
        with (
            patch("core.llm_tasks.async_task", return_value="tid"),
            patch("core.llm_tasks.result", return_value=None),
        ):
            run_conditional_anlage2_check(pf.pk)

        self.assertFalse(
            AnlagenFunktionsMetadaten.objects.filter(anlage_datei=pf).exists()
        )
        self.assertTrue(
            AnlagenFunktionsMetadaten.objects.filter(
                anlage_datei=other_pf
            ).exists()
        )

    def test_ajax_reset_all_reviews_resets_manual_fields(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"x"),
        )
        func = self.func
        res = AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf,
            funktion=func,
            is_negotiable_manual_override=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            quelle="parser",
            technisch_verfuegbar=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            quelle="ki",
            technisch_verfuegbar=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=func,
            quelle="manuell",
            technisch_verfuegbar=True,
        )
        self.client.login(username=self.superuser.username, password="pass")
        url = reverse("ajax_reset_all_reviews", args=[pf.pk])
        resp = self.client.post(url)
        self.assertEqual(resp.status_code, 200)
        res.refresh_from_db()
        self.assertFalse(
            FunktionsErgebnis.objects.filter(
                anlage_datei__project=projekt,
                anlage_datei=pf,
                funktion=func,
                quelle="manuell",
            ).exists()
        )
        self.assertIsNone(res.is_negotiable_manual_override)
        self.assertTrue(
            FunktionsErgebnis.objects.filter(
                anlage_datei__project=projekt,
                anlage_datei=pf,
                funktion=func,
                quelle="parser",
                technisch_verfuegbar=True,
            ).exists()
        )
        self.assertTrue(
            FunktionsErgebnis.objects.filter(
                anlage_datei__project=projekt,
                anlage_datei=pf,
                funktion=func,
                quelle="ki",
                technisch_verfuegbar=True,
            ).exists()
        )

    def test_hx_update_review_cell_toggles_manual_entry(self):
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("a.txt", b"x"),
        )
        func = self.func
        result = AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf,
            funktion=func,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            quelle="parser",
            technisch_verfuegbar=True,
        )
        FunktionsErgebnis.objects.create(
            anlage_datei=pf,
            funktion=self.func,
            quelle="ki",
            technisch_verfuegbar=False,
        )

        self.client.login(username=self.superuser.username, password="pass")
        url = reverse("hx_update_review_cell", args=[result.pk, "technisch_vorhanden"])

        resp = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(resp.status_code, 200)
        self.assertTrue(
            FunktionsErgebnis.objects.filter(
                anlage_datei=pf,
                funktion=func,
                quelle="manuell",
            ).exists()
        )
        # Manueller Eintrag liegt vor

        resp = self.client.post(url, HTTP_HX_REQUEST="true")
        self.assertEqual(resp.status_code, 200)
        self.assertFalse(
            FunktionsErgebnis.objects.filter(
                anlage_datei=pf,
                funktion=func,
                quelle="manuell",
            ).exists()
        )
        # Manueller Eintrag wurde entfernt


@pytest.mark.usefixtures("seed_db")
class GapReportTests(NoesisTestCase):
    def setUp(self):
        super().setUp()
        self.superuser = User.objects.get(username="frank")
        self.projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        self.pf1 = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("a.txt", b"data"),
            question_review={"1": {"hinweis": "Hinweis", "vorschlag": "V"}},
        )
        self.pf2 = BVProjectFile.objects.create(
            project=self.projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("b.txt", b"data"),
        )
        self.func = Anlage2Function.objects.create(name="Anmelden")
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=self.pf2,
            funktion=self.func,
            gap_summary="Extern",
            gap_notiz="Intern",
        )

    def test_tasks_return_text(self):
        self.client.login(username=self.superuser.username, password="pass")
        with patch("core.llm_tasks.query_llm", return_value="T1") as mock_q:
            text = summarize_anlage1_gaps(self.projekt)
        self.assertEqual(text, "T1")
        self.assertTrue(mock_q.called)

        with patch("core.llm_tasks.query_llm", return_value="T2") as mock_q:
            text = summarize_anlage2_gaps(self.projekt)
            prompt_sent = mock_q.call_args[0][0]
            self.assertIn("### Anmelden", prompt_sent.text)
            self.assertIn("- KI-Analyse:", prompt_sent.text)
            self.assertIn("- GAP-Anmerkung (Extern):", prompt_sent.text)
            self.assertIn(self.projekt.title, prompt_sent.text)
            self.assertNotIn("{gap_list}", prompt_sent.text)
            self.assertNotIn("{system_name}", prompt_sent.text)
            self.assertNotIn("{funktionen}", prompt_sent.text)
        self.assertEqual(text, "T2")
        self.assertTrue(mock_q.called)

    def test_view_saves_text(self):
        self.client.login(username=self.superuser.username, password="pass")
        url = reverse("gap_report_view", args=[self.projekt.pk])
        with patch("core.views.summarize_anlage1_gaps", return_value="A1"), patch(
            "core.views.summarize_anlage2_gaps", return_value="A2"
        ):
            resp = self.client.get(url)
            self.assertContains(resp, "A1")
            self.assertContains(resp, "A2")
            resp = self.client.post(url, {"text1": "E1", "text2": "E2"})
        self.assertRedirects(resp, reverse("projekt_detail", args=[self.projekt.pk]))
        self.pf1.refresh_from_db()
        self.pf2.refresh_from_db()
        self.assertEqual(self.pf1.gap_summary, "E1")
        self.assertEqual(self.pf2.gap_summary, "E2")

    def test_delete_gap_report(self):
        self.client.login(username=self.superuser.username, password="pass")
        self.pf1.gap_summary = "E1"
        self.pf2.gap_summary = "E2"
        self.pf1.save(update_fields=["gap_summary"])
        self.pf2.save(update_fields=["gap_summary"])
        url = reverse("delete_gap_report", args=[self.projekt.pk])
        resp = self.client.post(url)
        self.assertRedirects(resp, reverse("projekt_detail", args=[self.projekt.pk]))
        self.pf1.refresh_from_db()
        self.pf2.refresh_from_db()
        self.assertEqual(self.pf1.gap_summary, "")
        self.assertEqual(self.pf2.gap_summary, "")


@pytest.mark.usefixtures("seed_db")
class ProjektDetailGapTests(NoesisTestCase):
    def setUp(self):
        super().setUp()
        self.superuser = User.objects.get(username="frank")
        self.func = Anlage2Function.objects.create(name="Anmelden")

    def test_anlage1_gap_sets_flag(self):
        self.client.login(username=self.superuser.username, password="pass")
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=1,
            upload=SimpleUploadedFile("a.txt", b"data"),
            question_review={"1": {"vorschlag": "V"}},
        )
        resp = self.client.get(reverse("projekt_detail", args=[projekt.pk]))
        self.assertTrue(resp.context["can_gap_report"])

    def test_anlage2_gap_sets_flag(self):
        self.client.login(username=self.superuser.username, password="pass")
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf2 = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=2,
            upload=SimpleUploadedFile("b.txt", b"data"),
        )
        func = self.func
        AnlagenFunktionsMetadaten.objects.create(
            anlage_datei=pf2,
            funktion=func,
            supervisor_notes="Hinweis",
        )
        resp = self.client.get(reverse("projekt_detail", args=[projekt.pk]))
        self.assertTrue(resp.context["can_gap_report"])

    def test_anlage4_gap_sets_flag(self):
        self.client.login(username=self.superuser.username, password="pass")
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=4,
            upload=SimpleUploadedFile("c.txt", b"data"),
            manual_comment="Hinweis",
        )
        resp = self.client.get(reverse("projekt_detail", args=[projekt.pk]))
        self.assertTrue(resp.context["can_gap_report"])

    def test_anlage5_gap_sets_flag(self):
        self.client.login(username=self.superuser.username, password="pass")
        projekt = BVProject.objects.create(software_typen="A", beschreibung="x")
        pf5 = BVProjectFile.objects.create(
            project=projekt,
            anlage_nr=5,
            upload=SimpleUploadedFile("d.txt", b"data"),
        )
        ZweckKategorieA.objects.create(beschreibung="A")
        ZweckKategorieA.objects.create(beschreibung="B")
        Anlage5Review.objects.create(project_file=pf5)
        resp = self.client.get(reverse("projekt_detail", args=[projekt.pk]))
        self.assertTrue(resp.context["can_gap_report"])

class ManualGapDetectionTests(NoesisTestCase):
    """Tests für die Funktion ``_has_manual_gap``."""

    def test_detects_difference(self) -> None:
        """Erkennt eine Abweichung zwischen Dokument und manuellem Wert."""
        doc_data = {"technisch_vorhanden": True}
        manual_data = {"technisch_vorhanden": False}
        self.assertTrue(_has_manual_gap(doc_data, manual_data))

    def test_no_gap_when_equal(self) -> None:
        """Kein GAP, wenn Werte übereinstimmen."""
        doc_data = {"technisch_vorhanden": True}
        manual_data = {"technisch_vorhanden": True}
        self.assertFalse(_has_manual_gap(doc_data, manual_data))

    def test_gap_when_doc_missing(self) -> None:
        """Erkennt eine Lücke, wenn Dokumentdaten fehlen."""
        doc_data: dict = {}
        manual_data = {"technisch_vorhanden": True}
        self.assertTrue(_has_manual_gap(doc_data, manual_data))

    def test_no_gap_when_manual_missing(self) -> None:
        """Kein GAP, wenn manuelle Daten fehlen."""
        doc_data = {"technisch_vorhanden": True}
        manual_data = {"technisch_vorhanden": None}
        self.assertFalse(_has_manual_gap(doc_data, manual_data))

    def test_gap_with_additional_manual_field(self) -> None:
        """Erkennt eine Lücke bei zusätzlichen manuellen Feldern."""
        doc_data = {"technisch_vorhanden": True}
        manual_data = {
            "technisch_vorhanden": True,
            "neues_feld": False,
        }
        self.assertTrue(_has_manual_gap(doc_data, manual_data))

    def test_no_gap_when_parser_missing_special(self) -> None:
        """Kein GAP bei Spezialfeldern ohne Parser-Wert."""
        doc_data: dict = {}
        manual_data = {"einsatz_bei_telefonica": True}
        self.assertFalse(_has_manual_gap(doc_data, manual_data))

    def test_gap_only_on_difference_special(self) -> None:
        """GAP bei Spezialfeldern nur bei Abweichung."""
        doc_data = {"einsatz_bei_telefonica": False}
        manual_data = {"einsatz_bei_telefonica": True}
        self.assertTrue(_has_manual_gap(doc_data, manual_data))
        manual_data2 = {"einsatz_bei_telefonica": False}
        self.assertFalse(_has_manual_gap(doc_data, manual_data2))


class ResolveValueLogicTests(NoesisTestCase):
    """Tests für die Feldpriorisierung in ``_resolve_value``."""

    def test_doc_overrides_ai_for_special_fields(self) -> None:
        """Bei Spezialfeldern hat Parser-Vorrang vor KI."""
        val, src = _resolve_value(
            None, False, True, "einsatz_bei_telefonica", False, True
        )
        self.assertTrue(val)
        self.assertEqual(src, "Dokumenten-Analyse")

    def test_ai_still_used_for_regular_fields(self) -> None:
        """Bei normalen Feldern dominiert der KI-Wert."""
        val, _ = _resolve_value(
            None, False, True, "technisch_vorhanden", False, True
        )
        self.assertFalse(val)
